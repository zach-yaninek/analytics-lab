#!/usr/bin/env python3
"""End-to-end test of verify_with_db.py on synthetic fixtures with planted answers.

Builds three fake databases + workbooks under data/derived/fixture/ (gitignored, FAKE data only), runs the
verification on each with the fixture's own expected values, and asserts:
  * no step errors and no failed checks;
  * the combined clock verdict recovers the planted clock.
It also exercises the full-scan fallback (--no-chunk) and the index-only mode (--skip-scans).

It also runs db_followup.py on an extended fixture (the real DB's location and event-time columns) and checks the
last-10 code lists against an independent plain-pandas recomputation.

It runs fleet_db_extract.py (CLI) on the extended fixture with its fake installed-base list: once on the
fixture DB as built (machine ids, hosts and active days must equal db_followup.py's fu_devices.csv) and once on a
copy with planted gaps (months deleted, a silent device added), checking the schema, M ids against
common.build_maps, S ids against common.extend_maps, active months / first / last days against plain pandas, the
match summary against the planted list, and the privacy check (which must also catch planted leaks).

Finally it runs followups.py (CLI, docs/09) on the extended fixture and on a planted copy, checking the pairs, repeats,
request groups and Part A counts against an independent plain-pandas oracle, the random-timing null's calibration
(the fixture's ticket times are uniform), planted same-Category repeats, monitoring texts and a Touchscreen excess,
privacy and byte-identical reruns; and the follow-ups unit tests (test_followups_units.py).

Run:  cd /home/user/analytics-lab && python3 -I verification/tests/run_fixture_tests.py
"""
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))
sys.path.insert(0, str(HERE))

import contextlib  # noqa: E402
import datetime as dt  # noqa: E402
import hashlib  # noqa: E402
import io  # noqa: E402

import json  # noqa: E402
import shutil  # noqa: E402
import sqlite3  # noqa: E402
import subprocess  # noqa: E402

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

import common  # noqa: E402
import db_followup  # noqa: E402
import fleet_db_extract  # noqa: E402
import make_fixture  # noqa: E402
import verify_with_db  # noqa: E402

BASE = common.DERIVED / 'fixture'
VARIANTS = [  # name, clock, ingestion column, indexes, seed, extra run options
    ('local', 'local', 'ingested_at_utc', 'full', 7, {}),
    ('utc', 'utc', 'received_at', 'none', 11, {}),
    ('local_noingest', 'local', '', 'none', 23, {}),
    ('local_nochunk', 'local', 'ingested_at_utc', 'full', 7, {'no_chunk': True}),
]


def run_variant(name, clock, ingest, indexes, seed, opts):
    out = BASE / name
    db, xlsx, truth = make_fixture.build(out, clock, ingest, indexes, seed=seed)
    with contextlib.redirect_stdout(io.StringIO()) as log:
        s = verify_with_db.run(db, xlsx, out / 'out', out / 'private', expected_json=truth, chunk_rows=100_000,
                               **opts)
    (out / 'run.log').write_text(log.getvalue())
    errors = [k for k, v in s['steps'].items() if v['status'] == 'ERROR']
    fails = [c['check'] for c in s['checks'] if c['status'] == 'FAIL']
    verdict = s['decisions'].get('clock_verdict', {}).get('hypothesis')
    planted = 'machine_local_dst' if clock == 'local' else 'utc'
    ok = not errors and not fails and verdict == planted
    print(f"{name:15s} checks {s['check_counts']}  verdict={verdict}  planted={planted}  "
          f"{'OK' if ok else 'PROBLEM'}" + (f'  errors={errors}' if errors else '') + (f'  fails={fails}' if fails
                                                                                    else ''))
    return ok


def independent_last10(db, xlsx, case_ids):
    """Plain-pandas recomputation of the corrected-clock last-10 lists (no db_followup code)."""
    raw = common.load_raw(xlsx)
    cmap = common.build_maps(raw)['case']
    t = pd.DataFrame({'case_id': raw['Case Number'].map(common.clean).map(cmap),
                      'serial': raw['Serial #'].map(common.clean),
                      'opened': pd.to_datetime(raw['Date/Time Opened'])})
    t = t[t.case_id.isin(case_ids)]
    t['anchor'] = (t.opened.dt.tz_localize('America/New_York').dt.tz_convert('UTC').dt.tz_localize(None))
    con = sqlite3.connect(db)
    out = {}
    for serial, g in t.groupby('serial'):
        ev = pd.read_sql_query('SELECT transaction_time_utc AS tu, transaction_code AS code FROM transaction_events '
                               'WHERE external_device_id=?', con, params=(serial,))
        ev['tu'] = pd.to_datetime(ev.tu)
        ev = ev.dropna(subset=['tu']).sort_values(['tu', 'code'], ascending=[True, False])
        for r in g.itertuples():
            w = ev[(ev.tu >= r.anchor - pd.Timedelta(days=30)) & (ev.tu < r.anchor) & (ev.code != '00')].iloc[::-1]
            out[r.case_id] = ('|'.join(w.code.head(10)), '|'.join(pd.unique(w.code)[:10]))
    con.close()
    return out


def run_followup():
    """db_followup.py on an extended fixture (real DB's extra columns) with planted answers."""
    out = BASE / 'extended'
    db, xlsx, truth = make_fixture.build(out, 'utc', '', 'full', seed=11, extended=True)
    truth = json.loads(truth.read_text())
    con = sqlite3.connect(db)                       # privacy trap: free text typed into the state field
    con.execute("UPDATE devices SET state='Testville01, MA 02115' WHERE external_device_id='FX100075'")
    con.commit()
    con.close()
    ids = sorted(common.build_maps(common.load_raw(xlsx))['case'].values())[::9][:30]
    with contextlib.redirect_stdout(io.StringIO()) as log:
        s = db_followup.run(db, xlsx, out / 'fu_out', out / 'fu_private', pilot_ids=ids)
    (out / 'fu_run.log').write_text(log.getvalue())
    errors = [k for k, v in s['steps'].items() if v['status'] != 'done']
    fails = [c['check'] for c in s['checks'] if c['status'] == 'FAIL']
    f = s['facts']
    last10 = pd.read_csv(out / 'fu_out' / 'fu_ticket_last10.csv', dtype=str, keep_default_na=False)
    main = last10[last10.window == 'corrected_guard0']
    ind = independent_last10(db, xlsx, set(main.case_id))
    agree = sum(ind.get(r.case_id) == (r.last10_events, r.last10_distinct) for r in main.itertuples())
    problems = []
    if errors or fails:
        problems.append(f'errors={errors} fails={fails}')
    if len(main) != truth['linkage']['matched_cases']:
        problems.append(f"tickets {len(main)} != matched {truth['linkage']['matched_cases']}")
    if agree != len(main):
        problems.append(f'last-10 lists agree with independent pandas for {agree}/{len(main)} tickets')
    if f['F4']['ticketed devices: location_name vs ticket account'] != truth['extended']['name_match_ticketed_devices']:
        problems.append('location name-match counts differ from planted')
    if f['F5']['share of rows where event_time_offset == minus UTC offset of device zone'] < 0.99:
        problems.append('event_time_offset convention not recovered')
    leaks = [p.name for p in (out / 'fu_out').iterdir() if 'testville' in p.read_text(errors='ignore').lower()]
    if leaks:
        problems.append(f'free-text state leaked into public outputs: {leaks}')
    ok = not problems
    print(f"{'followup':15s} steps {[k for k in s['steps']]}  last-10 agree {agree}/{len(main)}  "
          f"{'OK' if ok else 'PROBLEM ' + '; '.join(problems)}")
    return ok


FXDB_SPEC = ['machine_id', 'serial_id', 'on_list', 'machine_type', 'host', 'shared_host', 'device_status',
             'first_active', 'last_active', 'active_months_in_period']
FXDB_GAPS = [  # planted on a copy of the extended DB: (device, SQL condition on transaction_date_local to delete)
    ('FX100010', "transaction_date_local >= '2025-10-01' AND transaction_date_local < '2026-01-01'"),  # 9 months
    ('FX100011', "transaction_date_local >= '2026-03-01'"),                                          # 7 months
    ('FX100012', "transaction_date_local < '2026-08-01'"),                     # only after the period: 0 months
    ('FX100013', "transaction_date_local < '2026-07-15'"),                     # unlisted; 1 month (2026-07)
]
SILENT_DEVICE = 'FX100090'      # added to `devices` with no events: no first/last day, 0 months


def run_extract(db, xlsx, equipment, out, fu_devices=''):
    """fleet_db_extract.py through its command line; returns (exit status, public frames)."""
    cmd = [sys.executable, '-I', str(HERE.parent / 'fleet_db_extract.py'), '--db', str(db), '--xlsx', str(xlsx),
           '--equipment', str(equipment), '--out', str(out / 'fxdb_out'), '--private', str(out / 'fxdb_private'),
           '--fu-devices', str(fu_devices)]
    p = subprocess.run(cmd, capture_output=True, text=True)
    (out / 'fxdb_run.log').write_text(p.stdout + p.stderr)
    read = {n: pd.read_csv(out / 'fxdb_out' / f'fxdb_{n}.csv', dtype=str, keep_default_na=False)
            for n in ('devices', 'match_summary', 'model_type_crosswalk')}
    return p.returncode, read, p.stdout


def independent_activity(db):
    """Plain pandas over every event: first/last day overall and in 2025-08..2026-07, distinct months in it."""
    con = sqlite3.connect(db)
    ev = pd.read_sql_query('SELECT external_device_id AS dev, transaction_date_local AS day FROM transaction_events '
                           'WHERE transaction_date_local IS NOT NULL', con)
    con.close()
    inp = ev[(ev.day >= '2025-08-01') & (ev.day <= '2026-07-31')]
    return pd.DataFrame({'first_active': ev.groupby('dev').day.min(), 'last_active': ev.groupby('dev').day.max(),
                         'first_active_in_period': inp.groupby('dev').day.min(),
                         'last_active_in_period': inp.groupby('dev').day.max(),
                         'active_months_in_period': inp.assign(m=inp.day.str[:7]).groupby('dev').m.nunique()})


def check_extract(db, xlsx, equipment, read, truth):
    """Problems in one extract run, against independent recomputations from the fixture files."""
    problems = []
    dv, ms, cw = read['devices'], read['match_summary'], read['model_type_crosswalk']
    if list(dv.columns) != fleet_db_extract.COLUMNS or list(dv.columns[:len(FXDB_SPEC)]) != FXDB_SPEC:
        problems.append(f'fxdb_devices columns {list(dv.columns)}')
    if list(ms.columns) != ['measure', 'value'] or list(cw.columns) != ['model', 'machine_type', 'machines']:
        problems.append('match summary / crosswalk columns')
    con = sqlite3.connect(db)
    dev = pd.read_sql_query('SELECT external_device_id, device_status, machine_type, clearview_url FROM devices', con)
    con.close()
    raw = common.load_raw(xlsx)
    eq = common.load_equipment(equipment)
    mmap = common.build_maps(raw, dev.external_device_id)['serial']
    smap = common.extend_maps(common.build_maps(raw), eq)['serial']
    inv = {v: k for k, v in mmap.items()}
    dv = dv.assign(device=dv.machine_id.map(inv))
    if len(dv) != len(dev) or set(dv.device) != set(dev.external_device_id.map(common.clean)):
        problems.append('M ids differ from common.build_maps(raw, device ids)')
    dev['device'] = dev.external_device_id.map(common.clean)
    x = dv.merge(dev, on='device', suffixes=('', '_db'))
    listed = set(eq.serial)
    want_sid = x.device.map(lambda d: smap[d] if d in listed else '')
    if not (x.serial_id == want_sid).all() or not x.serial_id[x.serial_id != ''].str.fullmatch(r'S\d{5}').all():
        problems.append('S ids differ from common.extend_maps(common.build_maps(raw), equipment)')
    if sorted(x.device[x.on_list == '1']) != sorted(truth['equipment']['listed_device_ids']):
        problems.append('on_list differs from the planted list')
    want_type = x.machine_type_db.map(lambda t: common.clean(t).lower() if common.clean(t) else 'Unknown')
    if not (x.machine_type == want_type).all():
        problems.append('machine_type normalisation')
    urls = sorted(set(dev.clearview_url.dropna()))
    want_host = x.clearview_url.map({u: f'Host {i:03d}' for i, u in enumerate(urls, 1)}).fillna('')
    want_shared = x.clearview_url.fillna('').str.rstrip('/').eq(common.SHARED_INTERNAL_HOST).astype(int).astype(str)
    if not (x.host == want_host).all() or not (x.shared_host == want_shared).all() or want_shared.eq('1').sum() == 0:
        problems.append('host / shared_host')
    act = independent_activity(db).reindex(x.device)
    act['active_months_in_period'] = act.active_months_in_period.fillna(0).astype(int).astype(str)
    for col in act.columns:
        got, want = x[col].to_numpy(), act[col].fillna('').astype(str).to_numpy()
        if not (got == want).all():
            problems.append(f'{col}: {int((got != want).sum())} devices differ from plain pandas')
    s = dict(zip(ms.measure, ms.value.astype(int)))
    e = truth['equipment']
    for k, v in {'devices': len(dev), 'devices on list': e['devices_on_list'], 'listed machines': e['listed_machines'],
                 'listed machines networked': e['devices_on_list'],
                 'listed machines not networked': e['listed_machines'] - e['devices_on_list'],
                 'leading-zero extra matches: devices (report only, not joined)': e['leading_zero_extra_devices'],
                 'leading-zero extra matches: listed machines not otherwise networked':
                     e['leading_zero_extra_devices'],
                 'active machine-months in period (all devices)': int(act.active_months_in_period.astype(int).sum())
                 }.items():
        if s.get(k) != v:
            problems.append(f'match summary {k!r}: {s.get(k)} != {v}')
    fam = eq.groupby('family').size()
    if any(s.get(f'family {f}: listed machines') != int(n) for f, n in fam.items()):
        problems.append('family counts')
    on = x[x.on_list == '1'].merge(eq[['serial', 'model']], left_on='device', right_on='serial')
    want_cw = on.groupby(['model', 'machine_type']).size().rename('machines').reset_index()
    got_cw = cw.assign(machines=cw.machines.astype(int))
    if not got_cw.equals(want_cw.sort_values(['model', 'machine_type']).reset_index(drop=True)):
        problems.append('model-type crosswalk')
    return problems


def run_db_extract():
    """fleet_db_extract.py on the extended fixture (after run_followup) and on a copy with planted gaps."""
    out = BASE / 'extended'
    db, xlsx, equipment = out / 'fake_telemetry.db', out / 'fake_cases.xlsx', out / make_fixture.EQUIPMENT_FILE
    truth = json.loads((out / 'fixture_truth.json').read_text())
    problems = []
    # 1. the DB db_followup.py ran on: same machine ids, hosts and first/last days as its fu_devices.csv
    status, read, log = run_extract(db, xlsx, equipment, out / 'fxdb_plain', out / 'fu_out' / 'fu_devices.csv')
    if status != 0:
        problems.append(f'plain run exit status {status}')
    for name in ('machine ids equal fu_devices.csv', 'host per machine equals fu_devices.csv',
                 'first_active per machine equals fu_devices.csv', 'last_active per machine equals fu_devices.csv'):
        if f'[PASS] {name}' not in log:
            problems.append(f'not PASS: {name}')
    problems += [f'plain: {p}' for p in check_extract(db, xlsx, equipment, read, truth)]
    # 2. a copy with planted gaps and a silent device
    gdb = out / 'fxdb_gaps.db'
    shutil.copyfile(db, gdb)
    con = sqlite3.connect(gdb)
    for device, cond in FXDB_GAPS:
        con.execute(f'DELETE FROM transaction_events WHERE external_device_id=? AND {cond}', (device,))
    con.execute("INSERT INTO devices (external_device_id, device_status, machine_type, clearview_url) "
                "VALUES (?, 'Active', NULL, NULL)", (SILENT_DEVICE,))
    con.commit()
    con.close()
    status, read, log = run_extract(gdb, xlsx, equipment, out / 'fxdb_gaps')
    if status != 0:
        problems.append(f'gaps run exit status {status}')
    problems += [f'gaps: {p}' for p in check_extract(gdb, xlsx, equipment, read, truth)]
    months = dict(zip(read['devices'].machine_id, read['devices'].active_months_in_period))
    mmap = common.build_maps(common.load_raw(xlsx), [f'FX{100000 + i}' for i in range(91)])['serial']
    planted = {'FX100010': '9', 'FX100011': '7', 'FX100012': '0', 'FX100013': '1', SILENT_DEVICE: '0'}
    if {d: months.get(mmap[d]) for d in planted} != planted:
        problems.append('planted active months not recovered')
    # 3. privacy: the check passed, no fake real string in any public file, and the scanner catches planted leaks
    for sub in ('fxdb_plain', 'fxdb_gaps'):
        if '[PASS] public text columns' not in (out / sub / 'fxdb_run.log').read_text():
            problems.append(f'{sub}: privacy check did not pass')
        for f in (out / sub / 'fxdb_out').iterdir():
            text = f.read_text()
            if any(t in text for t in ('FX1', 'FX7', 'FX8', 'Fake ', 'fake-host', 'example.test', 'Testville')):
                problems.append(f'{sub}: real fixture string in {f.name}')
    raw, eq = common.load_raw(xlsx), common.load_equipment(equipment)
    con = sqlite3.connect(db)
    dev = pd.read_sql_query('SELECT external_device_id, clearview_url FROM devices', con)
    con.close()
    dev['device'] = dev.external_device_id.map(common.clean)
    real = fleet_db_extract.real_strings(raw, eq, dev)
    trap = pd.DataFrame({'a': ['M00001', 'FX100001', 'x'], 'b': ['Host 001', 'Fake Memorial Hospital 03', 'y'],
                         'c': ['scrubex-3', 'see fake-host-02.example.test', 'Unknown'], 'n': [1, 2, 3]})
    if set(fleet_db_extract.text_leaks(trap, real)) != {'a', 'b', 'c'}:
        problems.append('privacy scanner missed a planted serial / account / url')
    if fleet_db_extract.text_leaks(trap.iloc[[0]].assign(a='M00001', c='scrubex-3'), real):
        problems.append('privacy scanner flags pseudonyms')
    ok = not problems
    print(f"{'db_extract':15s} devices {len(read['devices'])}  on list {int((read['devices'].on_list == '1').sum())}  "
          f"{'OK' if ok else 'PROBLEM ' + '; '.join(problems)}")
    return ok


def run_flags():
    """flags.py --write (the association study's text-flag export) on the extended fixture workbook, plus the copied
    regexes: planted matches and non-matches, and equality with their sources."""
    import ast
    import flags
    problems = []
    planted = [(flags.code_ref, 'Error 41 at the user door', True), (flags.code_ref, 'code 76 again', True),
               (flags.code_ref, 'replaced 2 doors', False), (flags.monitoring_template, 'Monitoring - Daily poll', True),
               (flags.telemetry_mention, 'seen in clearView', True),
               (flags.telemetry_mention, 'please clarify what happened', False),
               (flags.wrong_machine_note, 'Customer called wrong unit', True),
               (flags.wrong_machine_note, 'wrong part shipped', False)]
    for fn, text, want in planted:
        if fn(text) != want:
            problems.append(f'{fn.__name__}({text!r}) != {want}')

    def source_pattern(path, name):
        for node in ast.walk(ast.parse(path.read_text())):
            if isinstance(node, ast.Assign) and any(getattr(t, 'id', None) == name for t in node.targets):
                v = node.value
                return v.args[0].value if isinstance(v, ast.Call) else v.value
    if source_pattern(HERE.parent / 'timestamp_semantics.py', 'CODE_REF_RE') != flags.CODE_REF_RE.pattern:
        problems.append('CODE_REF_RE differs from timestamp_semantics.py')
    if source_pattern(HERE.parent / 'customer_fleet_audit.py', 'ERROR_NOTE') != flags.ERROR_NOTE.pattern:
        problems.append('ERROR_NOTE differs from customer_fleet_audit.py')
    xlsx = BASE / 'extended' / 'fake_cases.xlsx'
    if not xlsx.exists():
        make_fixture.build(BASE / 'extended', 'utc', '', 'full', seed=11, extended=True)
    out = BASE / 'extended' / 'flags'
    run = subprocess.run([sys.executable, '-I', str(HERE.parent / 'flags.py'), '--write', '--xlsx', str(xlsx),
                          '--out', str(out)], capture_output=True, text=True)
    if run.returncode != 0:
        problems.append(f'flags.py --write exit {run.returncode}: {run.stderr[-300:]}')
    else:
        f = pd.read_csv(out / 'flags_by_case.csv', dtype=str, keep_default_na=False)
        meta = json.loads((out / 'flags_by_case.meta.json').read_text())
        n_cases = len(common.load_raw(xlsx))
        bool_cols = [c for c in f.columns if c not in ('case_id', 'closed_at_eastern')]
        expected = {f'{rule}_{field}' for rule in ('monitoring_template', 'telemetry_mention', 'code_ref')
                    for field in ('reported', 'resolution', 'administrative')} | \
            {'wrong_machine_note_resolution', 'wrong_machine_note_administrative', 'wrong_machine_note'}
        if set(bool_cols) != expected:
            problems.append(f'columns {sorted(set(bool_cols) ^ expected)}')
        if len(f) != n_cases or meta['rows'] != n_cases:
            problems.append(f'rows {len(f)} / meta {meta["rows"]} != cases {n_cases}')
        if not set(f[bool_cols].to_numpy().ravel()) <= {'True', 'False'}:
            problems.append('non-boolean flag values')
        if f.apply(lambda col: col.str.contains('Fake', case=False)).to_numpy().any():
            problems.append('fixture text leaked into the export')
        if meta['workbook_sha256'] != common.sha256_file(xlsx):
            problems.append('meta sha256 mismatch')
        if not f.case_id.str.fullmatch(r'T\d{5}').all():
            problems.append('case ids not pseudonymous')
    ok = not problems
    print(f"{'flags --write':15s} {len(planted)} planted cases, copied regexes, export on the fixture  "
          f"{'OK' if ok else 'PROBLEM ' + '; '.join(problems)}")
    return ok


FU_FAKE_STRINGS = ('FX1', 'FX7', 'FX8', 'Fake ', 'fake-host', 'Testville', 'example.test')
FU_CUTOFF = pd.Timestamp('2026-08-02')


def followups_oracle(xlsx, equipment):
    """Plain-pandas recomputation (no followups.py code) of docs/09's pairs, repeats, request groups and the Part A
    ticket counts by model and age group. Monitoring uses flags.monitoring_template, the shared rule."""
    import flags
    raw = common.load_raw(xlsx)
    maps = common.build_maps(raw)
    c = pd.DataFrame({'case_id': raw['Case Number'].map(common.clean).map(maps['case']),
                      'serial': raw['Serial #'].map(common.clean),
                      'opened': pd.to_datetime(raw['Date/Time Opened']),
                      'trouble_code': raw['Trouble Code'].map(common.clean),
                      'category': raw['Trouble Code Category'].map(common.clean),
                      'record_type': raw['Case Record Type'].map(common.clean),
                      'monitoring': raw['Trouble Reported'].map(flags.monitoring_template).astype(bool)})
    c['machine_id'] = c.serial.map(maps['serial'])
    planned = c.trouble_code.isin(['Preventative Maintenance', 'Installation']) | (c.record_type == 'Install Case')
    u = c[~planned & ~c.category.isin(['Cancelled Request', 'Duplicate Case'])].copy()
    ref = pd.read_csv(HERE.parent / 'reference' / 'request_types_v0.csv', dtype=str, keep_default_na=False)
    lut = {(a, b): g for a, b, g in zip(ref.trouble_code, ref.category, ref.group)}
    fixed = {'Electronics': 'Electronics', 'Network': 'Network', 'Software': 'Software',
             'Operator Error': 'Operator error'}

    def group(tc, cat, mon):
        if mon:
            return 'Monitoring'
        tc, cat = ('[blank]' if pd.isna(tc) else tc), ('[blank]' if pd.isna(cat) else cat)
        if (tc, cat) in lut:
            return lut[(tc, cat)]
        return fixed.get(tc, 'Not categorised' if cat == '[blank]' else 'Other')
    u['group'] = [group(a, b, m) for a, b, m in zip(u.trouble_code, u.category, u.monitoring)]
    u = u.sort_values(['machine_id', 'opened', 'case_id'])
    nxt = u.groupby('machine_id')
    u['next_case_id'] = nxt.case_id.shift(-1)
    u['next_category'] = nxt.category.shift(-1)
    u['next_group'] = nxt.group.shift(-1)
    u['gap_days'] = (nxt.opened.shift(-1) - u.opened).dt.total_seconds() / 86400
    e = u[u.opened < FU_CUTOFF].copy()
    e['repeat_30'] = e.gap_days.between(0, 30)
    e['repeat_7'] = e.gap_days.between(0, 7)
    e['same_category'] = e.repeat_30 & e.category.notna() & (e.category == e.next_category)
    # Part A: listed machines, exposure and age at the midpoint (fleet_exposure F1-F3, re-derived here)
    eq = common.load_equipment(equipment)
    eq = eq[~eq.service_plan.isin(['Vogeltronix'])]
    start0, end = pd.Timestamp('2025-08-01'), pd.Timestamp('2026-09-01')
    t = u[u.serial.isin(set(eq.serial)) & (u.opened >= start0) & (u.opened < end)]
    first = t.groupby('serial').opened.min().dt.normalize()
    m = eq.assign(first=eq.serial.map(first))
    m['start'] = m[['install_date', 'first']].min(axis=1).clip(lower=start0)
    m['months'] = (end - m.start).dt.days.clip(lower=0) / (365.25 / 12)
    mid = m.start + (end - m.start) / 2
    m['young'] = ((mid - m.install_date).dt.days / 365.25).clip(lower=0) < 4
    m = m[m.months > 0]
    t = t[t.serial.isin(set(m.serial))].assign(young=lambda d: d.serial.map(dict(zip(m.serial, m.young))))
    return e, m, t


def followups_compare(out, e, m, t):
    """Problems where followups.py's outputs differ from the oracle."""
    problems = []
    p = pd.read_csv(out / 'rp_pairs.csv', keep_default_na=False, na_values=[''])
    p = p.set_index('case_id').reindex(e.case_id)
    if p.index.isna().any() or len(p) != len(e) or p.repeat_30.isna().any():
        return [f'eligible tickets differ ({len(p)} vs oracle {len(e)})']
    for col in ('repeat_30', 'repeat_7', 'same_category'):
        if (p[col].astype(bool).to_numpy() != e[col].to_numpy()).any():
            problems.append(f'{col} differs from the oracle')
    if (p.next_case_id.fillna('').to_numpy() != e.next_case_id.fillna('').to_numpy()).any():
        problems.append('next ticket differs from the oracle')
    gd = np.abs(p.gap_days.to_numpy(float) - e.gap_days.to_numpy(float))
    if np.nanmax(gd) > 1e-6 or (np.isnan(p.gap_days.to_numpy(float)) != np.isnan(e.gap_days.to_numpy())).any():
        problems.append('gap days differ from the oracle')
    if (p.index_group.to_numpy() != e.group.to_numpy()).any() or \
            (p.next_group.fillna('').to_numpy() != e.next_group.fillna('').to_numpy()).any():
        problems.append('request groups differ from the oracle')
    sup = pd.read_csv(out / 'ag_support.csv')
    mine = sup.assign(young=sup.age_group == 'Under 4').set_index(['model', 'young'])
    ora_t = t.assign(model=t.serial.map(dict(zip(m.serial, m.model)))).groupby(['model', 'young']).size()
    ora_m = m.groupby(['model', 'young']).agg(machines=('serial', 'size'), months=('months', 'sum'))
    ora = ora_m.join(ora_t.rename('tickets')).fillna(0)
    ora = ora[ora.machines > 0]
    j = mine.join(ora, rsuffix='_o', how='outer')
    if j[['machines', 'machines_o', 'tickets', 'tickets_o']].isna().any().any() or \
            (j.machines != j.machines_o).any() or (j.tickets != j.tickets_o).any() or \
            (np.abs(j.months - j.months_o) > 1e-6).any():
        problems.append('Part A support table (machines, months, tickets by model x age) differs from the oracle')
    rates = pd.read_csv(out / 'ag_type_rates.csv')
    g = rates[rates.level == 'group'].set_index('type')
    og = t.groupby(['group', 'young']).size()
    for grp, row in g.iterrows():
        if row.under4_tickets != og.get((grp, True), 0) or row.plus4_tickets != og.get((grp, False), 0):
            problems.append(f'Part A {grp} tickets by age differ from the oracle')
            break
    return problems


def plant_followups(src, dst, e, equipment):
    """A copy of the fixture workbook with known answers: 5 same-Category follow-ups 2 days after quiet index
    tickets, 4 follow-up tickets turned into monitoring templates, and 12 Touchscreen tickets on young LVRR-WP
    machines. Returns the planted case numbers."""
    wb = openpyxl.load_workbook(src)
    ws = wb[common.SHEET]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True)]
    raw = common.load_raw(src)
    cmap = common.build_maps(raw)['case']
    inv = {v: k for k, v in cmap.items()}
    by_case = {common.clean(r[col['Case Number']]): r for r in rows}
    opened = pd.to_datetime(raw['Date/Time Opened'])
    serials = raw['Serial #'].map(common.clean)
    quiet = []
    for r in e[e.category.notna() & (e.category != 'No Trouble Found') & (e.opened < '2026-06-01')].itertuples():
        later = opened[(serials == r.serial) & (opened > r.opened)]
        if later.empty or (later.min() - r.opened).days > 60:
            quiet.append(r)
    repeats = sorted(quiet, key=lambda x: x.case_id)[:5]
    new_rows, planted = [], {'repeat_index': [], 'repeat_new': [], 'monitoring': [], 'touchscreen': []}
    for k, r in enumerate(repeats):
        row = list(by_case[inv[r.case_id]])
        row[col['Case Number']] = f'99{k:06d}'
        row[col['Date/Time Opened']] = r.opened.to_pydatetime() + dt.timedelta(days=2)
        row[col['Trouble Reported']] = 'Fake planted repeat'
        row[col['Trouble Closed Date/Time']] = r.opened.to_pydatetime() + dt.timedelta(days=3)
        new_rows.append(row)
        planted['repeat_index'].append(r.case_id)
        planted['repeat_new'].append(f'99{k:06d}')
    follow = sorted(e[e.repeat_30 & ~e.next_case_id.isin([x.case_id for x in repeats])].next_case_id)[:4]
    for cid in follow:
        by_case[inv[cid]][col['Trouble Reported']] = 'Monitoring - Daily poll'
        planted['monitoring'].append(cid)
    eq = common.load_equipment(equipment)
    young = sorted(eq[(eq.model == 'LVRR-WP') & (eq.install_date >= '2022-11-01')].serial)[:4]
    proto = next(r for r in rows if common.clean(r[col['Model #']]) == 'LVRR-WP')
    acct = dict(zip(raw['Serial #'].map(common.clean), raw['Account Name']))
    for i, sn in enumerate(young):
        for j, month in enumerate((1, 4, 7)):
            row = list(proto)
            cn = f'98{3 * i + j:06d}'
            row[col['Case Number']] = cn
            row[col['Serial #']] = sn
            row[col['Account Name']] = acct.get(sn, proto[col['Account Name']])
            row[col['Date/Time Opened']] = dt.datetime(2026, month, 10, 10, 5 * i)
            row[col['Trouble Closed Date/Time']] = dt.datetime(2026, month, 11, 10, 0)
            row[col['Status']] = 'Trouble Closed'
            row[col['Trouble Code']], row[col['Trouble Code Category']] = 'Electronics', 'Touchscreen'
            row[col['Trouble Code Sub Category']] = None
            row[col['Trouble Reported']] = 'Fake planted touchscreen'
            new_rows.append(row)
            planted['touchscreen'].append(cn)
    ws.delete_rows(2, ws.max_row)
    for row in [by_case[common.clean(r[col['Case Number']])] for r in rows] + new_rows:
        ws.append(row)
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return planted


def run_followups_cli(xlsx, equipment, fu_dir, out):
    run = subprocess.run([sys.executable, '-I', str(HERE.parent / 'followups.py'), '--xlsx', str(xlsx),
                          '--equipment', str(equipment), '--fu-dir', str(fu_dir), '--out', str(out), '--no-pins',
                          '--n-boot', '200', '--n-perm', '200', '--n-sim', '200'], capture_output=True, text=True)
    (out.parent / f'{out.name}.log').write_text(run.stdout + run.stderr)
    return run.returncode, run.stdout


def run_followups():
    """followups.py (docs/09) on the extended fixture and a planted copy, against an independent oracle."""
    base = BASE / 'extended'
    xlsx, equipment, fu_dir = base / 'fake_cases.xlsx', base / make_fixture.EQUIPMENT_FILE, base / 'fu_out'
    problems = []
    outs = [base / 'followups_out', base / 'followups_out2']
    for out in outs:
        status, log = run_followups_cli(xlsx, equipment, fu_dir, out)
        if status != 0 or 'checks pass' not in log:
            problems.append(f'{out.name}: exit {status}')
    if not problems:
        e, m, t = followups_oracle(xlsx, equipment)
        problems += followups_compare(outs[0], e, m, t)
        hashes = [{f.name: hashlib.sha256(f.read_bytes()).hexdigest() for f in sorted(o.iterdir())} for o in outs]
        if hashes[0] != hashes[1]:
            problems.append('two runs differ: ' + ', '.join(k for k in hashes[0] if hashes[0][k] != hashes[1].get(k)))
        null = pd.read_csv(outs[0] / 'rp_null.csv').set_index(['variant', 'window'])
        for v in ('own window, uniform', 'full period window, uniform'):
            oe = null.loc[(v, '30 days'), 'observed_over_expected']
            if not 0.7 < oe < 1.3:
                problems.append(f'null calibration: {v} O/E {oe:.2f} (uniform planted times: expect about 1)')
        for f in outs[0].iterdir():
            if any(x in f.read_text() for x in FU_FAKE_STRINGS):
                problems.append(f'fixture identifier in {f.name}')
        base_ts = pd.read_csv(outs[0] / 'ag_type_adjusted.csv').query("variant == 'primary' and type == 'Touchscreen'")
        # planted copy
        pdir = base / 'followups_planted'
        pxlsx = pdir / 'fake_cases.xlsx'
        planted = plant_followups(xlsx, pxlsx, e, equipment)
        status, log = run_followups_cli(pxlsx, equipment, fu_dir, pdir / 'out')
        if status != 0:
            problems.append(f'planted run exit {status}')
        else:
            pe, pm, pt = followups_oracle(pxlsx, equipment)
            problems += [f'planted: {x}' for x in followups_compare(pdir / 'out', pe, pm, pt)]
            cmap = common.build_maps(common.load_raw(pxlsx))['case']
            pairs = pd.read_csv(pdir / 'out' / 'rp_pairs.csv').set_index('case_id')
            for idx, new in zip(planted['repeat_index'], planted['repeat_new']):
                r = pairs.loc[idx]
                if not (r.repeat_30 and r.same_category and r.next_case_id == cmap[new] and abs(r.gap_days - 2) < 1e-9):
                    problems.append(f'planted repeat after {idx} not recovered')
            mon = set(planted['monitoring'])
            inv = pairs[pairs.repeat_30 & (pairs.index.isin(mon) | pairs.next_case_id.isin(mon))]
            if len(inv) == 0 or (inv.pair_class != 'Monitoring involved').any():
                problems.append('planted monitoring follow-ups not classed "Monitoring involved"')
            if not (pairs.loc[pairs.index.isin(mon), 'index_group'] == 'Monitoring').all():
                problems.append('planted monitoring tickets not in the Monitoring group')
            ts = pd.read_csv(pdir / 'out' / 'ag_type_adjusted.csv').query("variant == 'primary' and "
                                                                         "type == 'Touchscreen'")
            if ts.empty or not (ts.share_ratio_mh.iloc[0] > 1 and ts.share_ratio_mh.iloc[0] >
                                base_ts.share_ratio_mh.iloc[0]):
                problems.append('planted Touchscreen excess not seen in the share ratio')
    unit = subprocess.run([sys.executable, '-I', str(HERE / 'test_followups_units.py')], capture_output=True,
                          text=True)
    if unit.returncode != 0:
        problems.append('unit tests: ' + unit.stdout.strip().splitlines()[-1])
    ok = not problems
    print(f"{'followups':15s} fixture + planted copy vs oracle, null calibration, privacy, reruns, unit tests  "
          f"{'OK' if ok else 'PROBLEM ' + '; '.join(problems)}")
    return ok


def main():
    results = [run_variant(*v) for v in VARIANTS]
    results.append(run_followup())
    results.append(run_db_extract())
    results.append(run_flags())
    results.append(run_followups())
    # index-only mode: steps 0-3, the pilot and the indexed DST-transition test must run without errors
    out = BASE / 'local'
    with contextlib.redirect_stdout(io.StringIO()):
        s = verify_with_db.run(out / 'fake_telemetry.db', out / 'fake_cases.xlsx', out / 'out_skip',
                               out / 'private_skip', skip_scans=True, expected_json=out / 'fixture_truth.json')
    errors = [k for k, v in s['steps'].items() if v['status'] == 'ERROR']
    ran = [k for k, v in s['steps'].items() if v['status'] == 'done']
    ok = not errors and {'0', '1', '2', '3', '4b', '6'} <= set(ran)
    print(f"{'skip_scans':15s} steps done {ran}  checks {s['check_counts']}  {'OK' if ok else 'PROBLEM'}")
    results.append(ok)
    print('ALL OK' if all(results) else 'SOME PROBLEMS')
    return 0 if all(results) else 1


if __name__ == '__main__':
    sys.exit(main())
