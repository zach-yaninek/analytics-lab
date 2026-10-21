#!/usr/bin/env python3
"""Build a small synthetic telemetry DB, a fake case workbook and a fake installed-base list for testing
verify_with_db.py, db_followup.py and fleet_db_extract.py.

Everything here is FAKE: account names, cities, serials, case numbers, hosts, owners and texts are invented.
The schema mirrors the real sources:
  * SQLite: transaction_events(external_device_id, transaction_date_local, transaction_time_local,
            transaction_time_utc, transaction_code[, <ingestion column>]),
            devices(external_device_id, device_status, machine_type, clearview_url),
            transaction_codes(transaction_code, transaction_description)
  * Workbook: sheet 'ipa_repair cases' with the 19 real column headers, plus a 'Summary' sheet.
  * Equipment list (fake_equipment.xlsx): sheet common.EQUIP_SHEET with the 22 real column headers. It lists 80 of
    the 90 devices, 5 ticket-only serials and 13 serials seen nowhere else (one padded with spaces), plus 2 serials
    that match a device only after stripping leading zeros. fixture_truth.json['equipment'] holds the planted
    counts.

The clock behaviour is planted so the verification's clock tests have a known answer:
  --clock local : transaction_time_local is each machine's own wall clock and observes DST where the hospital does,
                  so 2026-03-08 02:00-02:59 is empty and 2025-11-02 01:00-01:59 is doubled in DST-observing zones.
  --clock utc   : transaction_time_local actually holds UTC (no gap or doubled hour; zone profiles shift by
                  the UTC offset).
In both modes transaction_time_utc equals transaction_time_local except for a few planted rows (as in the
real DB, where only 155 rows differ). The ingestion column is always true UTC plus an upload delay.

fixture_truth.json lists values computed independently from the generated data (plain pandas). Pass it to
verify_with_db.py with --expected-json so every check compares against the planted answer.

Usage:
  python3 -I verification/tests/make_fixture.py --out data/derived/fixture/local --clock local
  python3 -I verification/tests/make_fixture.py --out data/derived/fixture/utc --clock utc \
      --ingest-col received_at --indexes none
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # verification/ (for common.py)

import argparse  # noqa: E402
import datetime as dt  # noqa: E402
import json  # noqa: E402
import sqlite3  # noqa: E402
from zoneinfo import ZoneInfo  # noqa: E402

import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402

import common  # noqa: E402

HEADERS = ['Date/Time Opened', 'Account Name', 'Hospital City, State, Zip', 'Shipping Country', 'Case Number',
           'Status', 'Dispatch #', 'Serial #', 'Model #', 'Product Name', 'Case Record Type', 'Case Owner',
           'Trouble Reported', 'Administrative Notes', 'Trouble Code', 'Trouble Code Category',
           'Trouble Code Sub Category', 'Trouble Resolution', 'Trouble Closed Date/Time']

# Fake locations: real state/province codes (so the zone parser works), fake city names, and zip / postal
# codes checked not to equal any real customer's (synthetic).
ZONES = {
    'Eastern': ('America/New_York', [('NY', '10098', 'USA'), ('MA', '02198', 'USA'), ('OH', '43298', 'United States'),
                                     ('ON', 'K0A 9Z9', 'Canada')]),
    'Central': ('America/Chicago', [('IL', '60698', 'USA'), ('TX', '75398', 'USA')]),
    'Mountain': ('America/Denver', [('CO', '80098', 'USA'), ('UT', '84198', 'USA')]),
    'Arizona (no DST)': ('America/Phoenix', [('AZ', '85098', 'USA')]),
    'Pacific': ('America/Los_Angeles', [('CA', '94198', 'USA'), ('WA', '98198', 'USA')]),
}
ACCOUNT_ZONES = (['Eastern'] * 10 + ['Central'] * 7 + ['Mountain'] * 4 + ['Arizona (no DST)'] * 3
                 + ['Pacific'] * 6)
MODELS = [('LVD-Remote-2.0', 'scrubEx', 'scrubEx / SD Case', ['scrubex2.0-2', 'scrubex2.0-3']),
          ('MV-2.0', 'scrubEx', 'scrubEx / SD Case', ['scrubex2.0-4']),
          ('alEx LC-D', 'alEx', 'alEx Case', ['alex2.0-1', 'alex2.0-3']),
          ('LC-30', 'alEx', 'alEx Case', ['mlc-1']),
          ('LVRR-WP', 'scrubEx', 'scrubEx / SD Case', ['scrubex-3'])]
ISSUES = [('Electro/Mechanical', 'Deposit Door', 'Lock Adjustment', 30),
          ('Electronics', 'Computer Board', 'Replacement', 15),
          ('Electronics', 'Touchscreen', None, 12),
          ('Miscellaneous', 'No Trouble Found', None, 8),
          ('Network', 'Connectivity', 'Training', 8),
          ('Miscellaneous', 'Cancelled Request', None, 4),
          ('Miscellaneous', 'Duplicate Case', None, 2),
          ('Preventative Maintenance', None, None, 2)]
# Hospital-like hour-of-day activity profile (local wall clock): shift changes near 07:00 and 19:00.
PROFILE = np.array([1.2, 1.0, 0.9, 0.9, 1.0, 1.5, 4.0, 6.5, 5.5, 4.5, 4.0, 3.8,
                    3.8, 3.6, 3.5, 3.8, 4.2, 4.5, 5.5, 6.2, 4.5, 3.0, 2.0, 1.2])
PROFILE = PROFILE / PROFILE.sum()
# 48 codes from Zac's lookup plus one malformed code => 49 rows, like the real table.
CODES = ['00', '01', '02', '03', '04', '05', '06', '07', '08', '09', '10', '11', '12', '13', '14', '15', '16', '17',
         '18', '19', '20', '21', '22', '23', '24', '25', '26', '27', '30', '31', '32', '33', '34', '35', '41', '43',
         '50', '51', '52', '53', '55', '60', '61', '66', '77', '79', '84', '90', 'X9']
FIRST_DATE, LAST_DATE = '2025-08-01', '2026-08-28'
# Installed-base list: the real sheet's 22 headers (common.load_equipment reads EQUIP_KEEP and EQUIP_DATES;
# common.equipment_real_strings reads EQUIP_PRIVATE_COLUMNS).
EQUIP_HEADERS = ['Serial Number', 'Service Status', 'Account Name', 'Hospital Address', 'Hospital City',
                 'Hospital State', 'Hospital Zip Code', 'Shipping Country', 'Phone', 'Product Family', 'Model',
                 'Manufacturing Date Code', 'Poll Position', 'Location', 'Last PM date', 'Equipment Ownership',
                 'Vendor Credentialing', 'Vendor Credentialing Notes', 'Service Plan', 'Current SC Expiration',
                 'Install Date', 'Go Live Date']
assert set(common.EQUIP_KEEP) | set(common.EQUIP_DATES.values()) | set(common.EQUIP_PRIVATE_COLUMNS) <= set(
    EQUIP_HEADERS)
EQUIPMENT_FILE = 'fake_equipment.xlsx'
UNLISTED_DEVICES = range(4, 90, 9)              # devices 4, 13, ..., 85 are not on the list (10 devices)
LEADING_ZERO_DEVICES = {76: '0', 85: '00'}      # ... but these two are listed with leading zeros


def write_equipment(out, devices, accounts, ticket_only_serials, seed):
    """Fake installed-base list. Uses its own random stream so the DB and workbook are unchanged."""
    rng = np.random.default_rng(seed + 1000)
    rows = []

    def row(serial, account, k, model, family, zone_state):
        st, zp, country = zone_state
        install = dt.datetime(2014, 1, 1) + dt.timedelta(days=int(rng.integers(0, 4500)))
        rows.append({'Serial Number': serial,
                     'Service Status': rng.choice(['Service Contract', 'Warranty', 'Time and Materials']),
                     'Account Name': account, 'Hospital Address': f'{100 + k} Fake Street',
                     'Hospital City': f'Testville{k:02d}', 'Hospital State': st, 'Hospital Zip Code': zp,
                     'Shipping Country': country, 'Phone': None if k % 7 == 0 else f'555-01{k:02d}',
                     'Product Family': family, 'Model': model, 'Manufacturing Date Code': f'FAKE-{2010 + k % 12}',
                     'Poll Position': f'{k % 4 + 1:02d}', 'Location': f'Fake Pharmacy {k % 3 + 1}',
                     'Last PM date': install + dt.timedelta(days=400),
                     'Equipment Ownership': rng.choice(['Customer Owned', 'Leased']),
                     'Vendor Credentialing': None, 'Vendor Credentialing Notes': None if k % 2 else 'Fake note',
                     'Service Plan': rng.choice(['Fake Plan A', 'Fake Plan B']),
                     'Current SC Expiration': install + dt.timedelta(days=3000),
                     'Install Date': install,
                     'Go Live Date': None if k % 5 == 0 else install + dt.timedelta(days=int(rng.integers(0, 60)))})

    for i, d in enumerate(devices):
        if i in UNLISTED_DEVICES and i not in LEADING_ZERO_DEVICES:
            continue
        serial = LEADING_ZERO_DEVICES.get(i, '') + d['id']
        if d['acct'] is not None:
            a = accounts[d['acct']]
            name = a['name'] if i % 6 else f"Fake Regional Health {d['acct']:02d}"   # list account != ticket account
            st = a['loc'].split(',')[1].split()
            loc = (st[0], ' '.join(st[1:]), a['country'])
        else:
            name, loc = f'Fake Community Clinic {i - 72:02d}', ZONES[d['zone']][1][0]
        row(serial, name, i % 40, d['model'], d['product'], loc)
    for j, s in enumerate(ticket_only_serials):
        row(s, f'Fake Memorial Hospital {j * 3 % 30:02d}', 40 + j, 'MV-2.0', 'scrubEx', ('NY', '10098', 'USA'))
    for j in range(13):
        serial = f'FX7{j:05d}' if j < 12 else f' FX7{j:05d} '              # whitespace removed by clean()
        model, family = MODELS[j % len(MODELS)][:2]
        row(serial, f'Fake Surgery Center {j % 5:02d}', 50 + j, model, family, ('CA', '94198', 'USA'))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = common.EQUIP_SHEET
    ws.append(EQUIP_HEADERS)
    for r in rows:
        ws.append([r[h] for h in EQUIP_HEADERS])
    path = Path(out) / EQUIPMENT_FILE
    wb.save(path)
    listed_devices = [d['id'] for i, d in enumerate(devices) if i not in UNLISTED_DEVICES]
    return path, {'listed_machines': len(rows), 'devices_on_list': len(listed_devices),
                  'listed_device_ids': listed_devices, 'ticket_only_serials': len(ticket_only_serials),
                  'serials_on_list_only': 13 + len(LEADING_ZERO_DEVICES),
                  'leading_zero_extra_devices': len(LEADING_ZERO_DEVICES)}


def fmt(epoch_seconds):
    """int64 epoch seconds -> 'YYYY-MM-DD HH:MM:SS' strings."""
    return np.char.replace(np.datetime_as_string(np.asarray(epoch_seconds, dtype='datetime64[s]'), unit='s'), 'T', ' ')


def build(out, clock='local', ingest_col='ingested_at_utc', indexes='full', scale=2.0, seed=7, extended=False):
    """extended=True adds the real DB's extra columns: devices(location_name, city, state, country, ...) and
    transaction_events(event_time_utc, event_time_offset), with planted answers in fixture_truth.json['extended']."""
    rng = np.random.default_rng(seed)
    out = Path(out)
    out.mkdir(parents=True, exist_ok=True)

    # ---------------- accounts, hosts, devices ----------------
    accounts = []
    for k, zone in enumerate(ACCOUNT_ZONES):
        st, zp, country = ZONES[zone][1][k % len(ZONES[zone][1])]
        accounts.append({'name': f'Fake Memorial Hospital {k:02d}', 'loc': f'Testville{k:02d}, {st} {zp}',
                         'country': country, 'zone': zone})
    host = {h: f'https://fake-host-{h:02d}.example.test' for h in range(1, 25)}
    shared = common.SHARED_INTERNAL_HOST + '/'

    def home_host(k):
        if k < 10:
            return k + 1                      # hosts 1-10
        if k < 20:
            return (k - 10) // 2 + 1          # hosts 1-5 gain two more accounts each
        return k - 9                          # hosts 11-20, one account each

    devices = []
    for i in range(90):
        model, product, rtype, types = MODELS[i % len(MODELS)]
        mtype = types[(i // len(MODELS)) % len(types)]
        if i % 17 == 3:
            mtype = None                       # -> 'Unknown'
        elif i % 13 == 5:
            mtype = ' ' + mtype.upper() + ' '  # case/whitespace variant of the same type
        acct = i % 30 if i < 72 else None
        if i < 72:
            url = host[home_host(acct)]
            if acct < 4 and i >= 30:
                url = host[11 + acct]          # second host for accounts 0-3
            if i in (60, 61, 62):
                url = shared
            if i in (63, 64):
                url = None
            zone = accounts[acct]['zone']
        else:                                  # devices with no tickets
            if i < 84:
                url = host[21 + (i - 72) % 4]  # hosts 21-24: no tickets at all
            elif i < 88:
                url = host[(i - 84) + 6]
            else:
                url = None if i == 88 else shared
            zone = ACCOUNT_ZONES[(i * 7) % 30]
        devices.append({'id': f'FX{100000 + i}', 'status': ['Active', 'Active', 'Inactive', 'Decommissioned'][i % 4],
                        'type': mtype, 'url': url, 'acct': acct, 'zone': zone, 'model': model, 'product': product,
                        'rtype': rtype})
    dev = pd.DataFrame(devices)

    # ---------------- tickets ----------------
    t0 = dt.datetime(2025, 8, 1)
    days = (dt.datetime(2026, 8, 31) - t0).days + 1
    issue_w = np.array([w for *_, w in ISSUES], dtype=float)
    issue_w /= issue_w.sum()
    tickets = []

    def ticket(serial, acct, model, product, rtype):
        a = accounts[acct]
        issue = ISSUES[rng.choice(len(ISSUES), p=issue_w)]
        status = rng.choice(['Trouble Closed', 'Dispatch', 'Closed'], p=[0.92, 0.05, 0.03])
        code, cat, sub = (None, None, None) if status == 'Dispatch' else issue[:3]
        opened = t0 + dt.timedelta(days=int(rng.integers(days)), hours=int(rng.integers(7, 18)),
                                   minutes=int(rng.integers(60)))
        closed = None if status == 'Dispatch' else opened + dt.timedelta(hours=float(rng.uniform(1, 72)))
        if closed is not None:
            closed = closed.replace(second=0, microsecond=0)
        tickets.append({'Date/Time Opened': opened, 'Account Name': a['name'], 'Hospital City, State, Zip': a['loc'],
                        'Shipping Country': a['country'], 'Case Number': None, 'Status': status,
                        'Dispatch #': None if rng.random() < 0.12 else f'D-{int(rng.integers(10**6)):06d}',
                        'Serial #': serial, 'Model #': model, 'Product Name': product,
                        'Case Record Type': 'General Case' if cat == 'Cancelled Request' else rtype,
                        'Case Owner': f'Fake Owner {"ABC"[int(rng.integers(3))]}',
                        'Trouble Reported': f'Fake reported problem {int(rng.integers(100))}',
                        'Administrative Notes': None if rng.random() < 0.3 else 'Fake admin note',
                        'Trouble Code': code, 'Trouble Code Category': cat, 'Trouble Code Sub Category': sub,
                        'Trouble Resolution': None if rng.random() < 0.1 else 'Fake resolution',
                        'Trouble Closed Date/Time': closed})

    for d in devices[:72]:
        for _ in range(int(rng.integers(2, 7))):
            ticket(d['id'], d['acct'], d['model'], d['product'], d['rtype'])
    ticket(devices[7]['id'], 8, devices[7]['model'], devices[7]['product'], devices[7]['rtype'])  # serial, 2 accounts
    for j in range(10):                                                                            # no device match
        ticket(f'FX8{j:05d}', j * 3 % 30, 'MV-2.0', 'scrubEx', 'scrubEx / SD Case')
    ticket('fx100005', 5, devices[5]['model'], devices[5]['product'], devices[5]['rtype'])           # case-only diff
    ticket(' FX100006 ', 6, devices[6]['model'], devices[6]['product'], devices[6]['rtype'])         # whitespace
    ticket(devices[9]['id'], 9, devices[9]['model'], devices[9]['product'], 'Install Case')          # planned work
    order = rng.permutation(len(tickets))
    for n, idx in enumerate(order):
        tickets[idx]['Case Number'] = f'{10000000 + 37 * n:08d}'
    tickets[int(order[0])]['Trouble Closed Date/Time'] = '11/21/2025 7:11 AM'                        # text close time
    tickets.sort(key=lambda r: r['Date/Time Opened'])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = common.SHEET
    ws.append(HEADERS)
    for r in tickets:
        ws.append([r[h] for h in HEADERS])
    summ = wb.create_sheet('Summary')
    summ.append(['Top repair issue (FAKE)', 'Cases', 'Share'])
    for cat, n in pd.Series([r['Trouble Code Category'] for r in tickets]).value_counts().head(5).items():
        summ.append([cat, int(n), round(n / len(tickets), 3)])
    xlsx = out / 'fake_cases.xlsx'
    wb.save(xlsx)
    _, equipment_truth = write_equipment(out, devices, accounts, [f'FX8{j:05d}' for j in range(5)], seed)

    # ---------------- telemetry ----------------
    hours = pd.date_range('2025-07-31 12:00', '2026-08-29 12:00', freq='h', tz='UTC')
    h_epoch = (hours.tz_localize(None) - pd.Timestamp('1970-01-01')) // pd.Timedelta('1s')
    h_epoch = np.asarray(h_epoch, dtype=np.int64)
    offsets = {}
    for zone, (iana, _) in ZONES.items():
        tz = ZoneInfo(iana)
        offsets[zone] = np.array([int(ts.to_pydatetime().astimezone(tz).utcoffset().total_seconds()) for ts in hours],
                                 dtype=np.int64)
    ids = list(dev.id) + ['FX999999']                       # last one: transmitting device missing from `devices`
    zones = list(dev.zone) + ['Central']
    rates = np.append(rng.uniform(8, 22, size=len(dev)), 1.0) * scale
    dev_i, hr_i = [], []
    for k, zone in enumerate(zones):
        local_hour = ((h_epoch + offsets[zone]) // 3600) % 24
        counts = rng.poisson(rates[k] * PROFILE[local_hour])
        hr_i.append(np.repeat(np.arange(len(hours)), counts))
        dev_i.append(np.full(int(counts.sum()), k))
    dev_i, hr_i = np.concatenate(dev_i), np.concatenate(hr_i)
    utc = h_epoch[hr_i] + rng.integers(0, 3600, size=len(hr_i))
    zone_names = list(ZONES)
    off_matrix = np.vstack([offsets[z] for z in zone_names])               # zones x hours, seconds
    zone_idx = np.array([zone_names.index(z) for z in zones])
    local = utc + off_matrix[zone_idx[dev_i], hr_i]
    recorded = local if clock == 'local' else utc
    other = utc if clock == 'local' else local
    tl = fmt(recorded).astype(object)
    dl = np.array([s[:10] for s in tl], dtype=object)
    keep = (dl >= FIRST_DATE) & (dl <= LAST_DATE)
    dev_i, utc, other, tl, dl = dev_i[keep], utc[keep], other[keep], tl[keep], dl[keep]
    order = np.argsort(utc, kind='stable')                   # rows stored in arrival order, like a log
    dev_i, utc, other, tl, dl = dev_i[order], utc[order], other[order], tl[order], dl[order]
    n = len(tl)
    tu = tl.copy()
    plant = rng.choice(n, size=12, replace=False)
    tu[plant[:7]] = fmt(other[plant[:7]]).astype(object)    # 7 rows: the "utc" column holds the other clock
    tu[plant[7:9]] = None                                   # 2 rows: missing utc string
    dl[plant[9:12]] = None                                  # 3 rows: missing local date
    weights = np.array([0.93] + [0.07 / 47] * 47)
    code = np.array(CODES[:48], dtype=object)[rng.choice(48, size=n, p=weights / weights.sum())]
    code[rng.choice(n, size=3, replace=False)] = 'X9'       # malformed (in the code table, not in Zac's lookup)
    devid = np.array(ids, dtype=object)[dev_i]
    latency = np.where(rng.random(n) < 0.95, rng.exponential(120, n), rng.uniform(3600, 36 * 3600, n)).astype(np.int64)
    ingest = fmt(utc + latency).astype(object)

    db = out / 'fake_telemetry.db'
    if db.exists():
        db.unlink()
    con = sqlite3.connect(db)
    extra = f', {ingest_col} TEXT' if ingest_col else ''
    dev_extra = ', location_name TEXT, city TEXT, state TEXT, country TEXT' if extended else ''
    ev_extra = ', event_time_utc TEXT, event_time_offset REAL' if extended else ''
    con.executescript(f"""
        CREATE TABLE devices (external_device_id TEXT, device_status TEXT, machine_type TEXT, clearview_url TEXT,
                              last_seen_utc TEXT{dev_extra});
        CREATE TABLE transaction_codes (transaction_code TEXT, transaction_description TEXT);
        CREATE TABLE transaction_events (external_device_id TEXT, transaction_date_local TEXT,
            transaction_time_local TEXT, transaction_time_utc TEXT, transaction_code TEXT{extra}{ev_extra});
    """)
    planted_names = {}
    if extended:
        full = {'NY': 'New York', 'MA': 'Massachusetts', 'IL': 'Illinois', 'CA': 'California', 'CO': 'Colorado'}
        rows = []
        for i, d in enumerate(devices):
            acct = d['acct'] if d['acct'] is not None else (((i - 72) % 30) if i < 84 else None)
            a = accounts[acct] if acct is not None else None
            st = a['loc'].split(',')[1].split()[0] if d['acct'] is not None else ZONES[d['zone']][1][0][0]
            if d['acct'] is not None:
                kind = ['exact'] * 7 + ['normalised', 'different', 'missing']
                kind = kind[i % 10]
                name = {'exact': a['name'], 'normalised': f"The {a['name']}, Inc.",
                        'different': f"Fake Health System {d['zone']}", 'missing': None}[kind]
                planted_names[d['id']] = kind
            else:
                name = a['name'] if a else None
            rows.append((d['id'], d['status'], d['type'], d['url'], '2026-08-28 00:00:00', name,
                         a['loc'].split(',')[0] if a else None, full.get(st, st) if i % 5 == 0 else st,
                         a['country'] if a else 'USA'))
        con.executemany('INSERT INTO devices VALUES (?,?,?,?,?,?,?,?,?)', rows)
    else:
        con.executemany('INSERT INTO devices VALUES (?,?,?,?,?)',
                        [(d['id'], d['status'], d['type'], d['url'], '2026-08-28 00:00:00') for d in devices])
    con.executemany('INSERT INTO transaction_codes VALUES (?,?)', [(c, f'Fake description {c}') for c in CODES])
    cols = [devid, dl, tl, tu, code] + ([ingest] if ingest_col else [])
    if extended:
        hour_of_row = np.searchsorted(h_epoch, utc, side='right') - 1
        dev_zone_idx = zone_idx[dev_i]
        offset_h = -off_matrix[dev_zone_idx, np.clip(hour_of_row, 0, len(h_epoch) - 1)] / 3600.0
        cols += [fmt(utc + np.minimum(latency, 600)).astype(object), offset_h.astype(float)]
    con.executemany(f'INSERT INTO transaction_events VALUES ({",".join("?" * len(cols))})', zip(*cols))
    if indexes == 'full':
        con.executescript("""
            CREATE INDEX ix_events_device_date ON transaction_events(external_device_id, transaction_date_local);
            CREATE INDEX ix_events_date ON transaction_events(transaction_date_local);
            CREATE INDEX ix_events_code ON transaction_events(transaction_code);
        """)
    con.commit()
    con.close()

    # ---------------- independent truth (plain pandas on the generated arrays) ----------------
    ev = pd.DataFrame({'dev': devid, 'dl': dl, 'tl': tl, 'tu': tu, 'code': code})
    t = pd.DataFrame(tickets)
    t['serial'] = t['Serial #'].str.strip()
    dev_ids = set(dev.id)
    dev_by_id = dev.set_index('id')
    known = dev_ids | set(t.serial)
    matched = t[t.serial.isin(dev_ids)].copy()
    norm_type = dev_by_id.type.map(lambda x: x.strip().lower() if isinstance(x, str) and x.strip() else 'Unknown')
    matched['mtype'] = matched.serial.map(norm_type)
    matched['host'] = matched.serial.map(dev_by_id.url)
    month = ev.dl.str[:7]
    dkey = ev.dl.fillna('<null>')
    opened = t['Date/Time Opened'].sort_values()
    differ = ev.tu.isna() | (ev.tu != ev.tl)

    hosts_with_tickets = matched.dropna(subset=['host']).groupby('host')['Account Name'].nunique()
    host_class = {}
    for url in dev.url.dropna().unique():
        if url.rstrip('/') == common.SHARED_INTERNAL_HOST:
            host_class[url] = 'Shared internal host'
        elif url not in hosts_with_tickets.index:
            host_class[url] = 'No tickets'
        else:
            host_class[url] = 'Single-account' if hosts_with_tickets[url] == 1 else 'Multi-account'
    truth = {
        '_note': 'Planted answers for the synthetic fixture; computed with plain pandas in make_fixture.py.',
        'devices_rows': len(dev), 'devices_unique': True, 'codes_rows': len(CODES),
        'record_quality': {'cases': len(t), 'distinct_cases': int(t['Case Number'].nunique()),
                           'customers': int(t['Account Name'].nunique()), 'case_machines': int(t.serial.nunique()),
                           'models': int(t['Model #'].nunique()),
                           'first_opened': opened.iloc[0].strftime('%Y-%m-%d %H:%M:%S'),
                           'last_opened': opened.iloc[-1].strftime('%Y-%m-%d %H:%M:%S')},
        'linkage': {'cases': len(t), 'matched_cases': len(matched), 'matched_machines': int(matched.serial.nunique()),
                    'matched_pct': round(100 * len(matched) / len(t), 2)},
        'mapping_summary': {
            'Case models spanning more than one telemetry type':
                int((matched.groupby('Model #').mtype.nunique() > 1).sum()),
            'Customers spanning more than one observed host':
                int((matched.groupby('Account Name').host.nunique() > 1).sum()),
            'Hosts linked to more than one case customer': int((hosts_with_tickets > 1).sum())},
        'host_classes': pd.Series(host_class).value_counts().to_dict(),
        'telemetry_groups': int(pd.DataFrame({'d': ev.dev, 'm': month.fillna('<null>'), 'c': ev.code})
                                .drop_duplicates().shape[0]),
        'telemetry_rows': n,
        'clocks_differ': int(differ.sum()),
        'missing_time': int((ev.tu.isna() | ev.dl.isna()).sum()),
        'null_local_date': int(ev.dl.isna().sum()),
        'date_time_mismatch': int((ev.tl.str[:10] != dkey).sum()),
        'unmapped_telemetry_devices': int(len(set(ev.dev) - known)),
        'coverage': {'first_transaction_date': ev.dl.dropna().min(), 'last_transaction_date': ev.dl.dropna().max(),
                     'represented_transactions': n, 'transacting_machines': int(len(set(ev.dev) & known))},
        'code_totals': {k: int(v) for k, v in ev.code.value_counts().sort_index().items()},
        'clock_differ_rows': int(differ.sum()),
        'clock_hypothesis': 'machine_local_dst' if clock == 'local' else 'utc',
        'dst_transition_hypothesis': 'machine_local_dst' if clock == 'local' else 'utc',
        'ingestion_hypothesis': ('machine_local_dst' if clock == 'local' else 'utc') if ingest_col else None,
        'planted': {'clock': clock, 'ingest_col': ingest_col or None, 'indexes': indexes, 'seed': seed,
                    'rows': n, 'devices': len(dev), 'tickets': len(t)},
    }
    truth['equipment'] = equipment_truth
    if extended:
        truth['extended'] = {'name_match_ticketed_devices': pd.Series(planted_names).value_counts().to_dict(),
                             'event_time_offset_convention': 'minus the UTC offset (hours west of UTC positive)',
                             'event_time_lag_max_seconds': 600}
    (out / 'fixture_truth.json').write_text(json.dumps(truth, indent=2, default=str))
    print(f'fixture written to {out}: {n:,} events, {len(dev)} devices, {len(t)} tickets, clock={clock}, '
          f'ingest_col={ingest_col or None}, indexes={indexes}')
    return db, xlsx, out / 'fixture_truth.json'


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--out', default=str(common.DERIVED / 'fixture' / 'local'))
    ap.add_argument('--clock', choices=['local', 'utc'], default='local')
    ap.add_argument('--ingest-col', default='ingested_at_utc', help="'' to omit the ingestion column")
    ap.add_argument('--indexes', choices=['full', 'none'], default='full')
    ap.add_argument('--scale', type=float, default=2.0, help='multiplies the event rate (2.0 ~ 1.1M rows)')
    ap.add_argument('--seed', type=int, default=7)
    ap.add_argument('--extended', action='store_true', help="add the real DB's location and event-time columns")
    a = ap.parse_args(argv)
    build(a.out, a.clock, a.ingest_col, a.indexes, a.scale, a.seed, a.extended)


if __name__ == '__main__':
    main()
