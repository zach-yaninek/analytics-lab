"""Adversarial logic audit of the case-side (Excel) parts of Zac's notebook.

Covers notebook cells 4, 8, 13, 15, 17, 18, 20, 29, 31, 32 and 37. Every number the audit
reports is computed here. Outputs are aggregate or pseudonymous only (Customer NNN, T#####,
S#####, generic labels). Free text is read in memory to set yes/no flags and is never written.

Run: cd /home/user/analytics-lab && python3 -I verification/audit_case_logic.py
Writes verification/outputs/audit_case_*.csv
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))

import difflib
import math
import re
import zipfile

import numpy as np
import openpyxl
import pandas as pd

from common import (CASE_VIEWS, OUTPUTS, XLS, ZAC_WORKBOOK_SHA256, build_cases, build_maps, case_database, clean,
                    load_raw, sha256_file)

OUTPUTS.mkdir(parents=True, exist_ok=True)
pd.set_option('display.width', 220)
pd.set_option('display.max_columns', 30)


def out(df, name):
    path = OUTPUTS / f'audit_case_{name}.csv'
    df.to_csv(path, index=False)
    print(f'  -> {path.relative_to(OUTPUTS.parents[1])} ({len(df)} rows)')


def banner(text):
    print('\n' + '=' * 100 + '\n' + text + '\n' + '=' * 100)


# --------------------------------------------------------------------------------------------
# Load exactly as cell 4 (see 01_excel_reproducibility.py for the pandas-3 NaN -> None fix).
# --------------------------------------------------------------------------------------------
raw0 = load_raw()
raw = raw0.astype(object).where(raw0.notna(), None)
maps = build_maps(raw)
cases = build_cases(raw, maps)
conn = case_database(cases)
sql = lambda s: pd.read_sql_query(s, conn)
assert not cases.isin(['nan', 'None', 'NaT']).any().any()

# Helper flags derived in memory (never exported as text).
reported = raw['Trouble Reported'].fillna('').astype(str)
# Proactive remote-monitoring tickets opened by IPA from clearView polling, e.g. 'Monitoring - Daily | ...'.
PROACTIVE_RE = r'monitoring\s*-?\s*(?:daily|poll|weekly)|^\s*monitoring\b'
cases['proactive_monitoring'] = reported.str.contains(PROACTIVE_RE, case=False, regex=True).astype(int)
cases['mentions_clearview_or_monitoring'] = reported.str.contains(r'monitoring|clear\s*view', case=False,
                                                                 regex=True).astype(int)
location = raw['Hospital City, State, Zip'].map(clean)
cases['state'] = location.str.extract(r',\s*([A-Z]{2})\b')[0]
cases.to_sql('cases', conn, index=False, if_exists='replace')  # views re-bind to the new table
conn.executescript(CASE_VIEWS)
base = sql('SELECT * FROM case_base')
service = sql('SELECT * FROM service_cases')
# Positional alignment with `raw` is used below for in-memory text flags; make sure it holds.
assert (base.case_id.to_numpy() == cases.case_id.to_numpy()).all()

# --------------------------------------------------------------------------------------------
banner('0. Baseline: reproduce the Zac numbers this audit perturbs')
# --------------------------------------------------------------------------------------------
REPEAT_VIEW_SQL = """
WITH ordered AS (
 SELECT *,LEAD(opened_at) OVER(PARTITION BY machine_id ORDER BY opened_at,case_id) next_opened,
 LEAD(category) OVER(PARTITION BY machine_id ORDER BY opened_at,case_id) next_category
 FROM {src})
SELECT *,julianday(next_opened)-julianday(opened_at) gap_days,
 CASE WHEN julianday(next_opened)-julianday(opened_at) BETWEEN 0 AND 30 THEN 1 ELSE 0 END repeat_30,
 CASE WHEN julianday(next_opened)-julianday(opened_at) BETWEEN 0 AND 30 AND category IS NOT NULL AND category=next_category THEN 1 ELSE 0 END same_category_30
FROM ordered WHERE opened_at<'2026-08-02'
"""


def repeat_frame(frame):
    """Zac's cell-31 repeat view, run verbatim on an arbitrary subset of service tickets."""
    frame = frame[[c for c in frame.columns if c not in ('next_opened', 'next_category', 'gap_days',
                                                          'repeat_30', 'same_category_30')]]
    frame.to_sql('subset_cases', conn, index=False, if_exists='replace')
    return sql(REPEAT_VIEW_SQL.format(src='subset_cases'))


def repeat_summary(frame):
    r = repeat_frame(frame)
    n = len(r)
    return {'unplanned_population': len(frame), 'eligible_tickets': n, 'machines': r.machine_id.nunique(),
            'repeat_tickets': int(r.repeat_30.sum()), 'repeat_pct': round(100 * r.repeat_30.sum() / n, 2),
            'same_category_repeats': int(r.same_category_30.sum()),
            'same_category_repeat_pct': round(100 * r.same_category_30.sum() / n, 2)}


baseline_rep = repeat_frame(service)
zac = {
    'service_group Other recorded issue': 6508, 'service_group No Trouble Found': 577,
    'service_group Cancelled request': 247, 'service_group Missing trouble code': 101,
    'service_group Planned work': 23, 'service_group Duplicate case': 15, 'unplanned population': 7186,
    'repeat eligible_tickets': 6665, 'repeat machines': 3495, 'repeat_tickets': 1574,
    'same_category_repeats': 515, 'simultaneous_followups': 0,
}
ours = {f'service_group {k}': int(v) for k, v in base.service_group.value_counts().items()}
ours.update({'unplanned population': len(service), 'repeat eligible_tickets': len(baseline_rep),
             'repeat machines': baseline_rep.machine_id.nunique(),
             'repeat_tickets': int(baseline_rep.repeat_30.sum()),
             'same_category_repeats': int(baseline_rep.same_category_30.sum()),
             'simultaneous_followups': int(((baseline_rep.repeat_30 == 1) & (baseline_rep.gap_days == 0)).sum())})
repro = pd.DataFrame([{'metric': k, 'zac_printed': v, 'recomputed': ours[k], 'match': v == ours[k]}
                      for k, v in zac.items()])
print(repro.to_string(index=False))
assert repro.match.all()
BASE = repeat_summary(service)
print(BASE)

# --------------------------------------------------------------------------------------------
banner('1. Cell 15 service groups: status, record types, label placement')
# --------------------------------------------------------------------------------------------
sg_status = pd.crosstab(base.service_group, base.status, margins=True, margins_name='All').reset_index()
print(sg_status.to_string(index=False))
out(sg_status, 'service_group_by_status')

status_month = pd.crosstab(base.opened_at.str[:7], base.status).reset_index().rename(columns={'opened_at': 'month'})
print(status_month.to_string(index=False))
out(status_month, 'status_by_month')

non_tc = base[base.status != 'Trouble Closed']
missing_code = base[base.service_group == 'Missing trouble code']
facts = {
    'missing_code_total': len(missing_code),
    'missing_code_status_Dispatch': int((missing_code.status == 'Dispatch').sum()),
    'missing_code_status_not_Trouble_Closed': int((missing_code.status != 'Trouble Closed').sum()),
    'missing_code_opened_2026_08': int(missing_code.opened_at.str.startswith('2026-08').sum()),
    'missing_code_among_Trouble_Closed': int((missing_code.status == 'Trouble Closed').sum()),
    'trouble_closed_total': int((base.status == 'Trouble Closed').sum()),
    'non_trouble_closed_total': len(non_tc),
    'non_trouble_closed_in_unplanned': int(non_tc.case_id.isin(service.case_id).sum()),
    'non_trouble_closed_opened_2026_08': int(non_tc.opened_at.str.startswith('2026-08').sum()),
    'non_trouble_closed_missing_close_time': int(non_tc.closed_at.isna().sum()),
    'trouble_closed_missing_close_time': int(base[base.status == 'Trouble Closed'].closed_at.isna().sum()),
    'aug2026_tickets': int(base.opened_at.str.startswith('2026-08').sum()),
    'aug2026_not_trouble_closed': int(non_tc.opened_at.str.startswith('2026-08').sum()),
    'Closed_status_months_distinct': int(base[base.status == 'Closed'].opened_at.str[:7].nunique()),
    'record_type_dot': int((base.record_type == '.').sum()),
    'record_type_dot_service_group': '; '.join(base[base.record_type == '.'].service_group.unique()),
    'install_case_total': int((base.record_type == 'Install Case').sum()),
    'install_case_with_repair_trouble_code': int(((base.record_type == 'Install Case') &
                                                   ~base.trouble_code.isin(['Installation', 'Preventative Maintenance'])).sum()),
    'installation_code_total': int((base.trouble_code == 'Installation').sum()),
    'installation_code_and_install_case': int(((base.trouble_code == 'Installation') &
                                                (base.record_type == 'Install Case')).sum()),
    'pm_code_total': int((base.trouble_code == 'Preventative Maintenance').sum()),
    'ntf_in_category': int((base.category == 'No Trouble Found').sum()),
    'ntf_trouble_codes': '; '.join(sorted(base[base.category == 'No Trouble Found'].trouble_code.dropna().unique())),
}
label_cols = {'trouble_code': base.trouble_code, 'category': base.category, 'subcategory': base.subcategory,
              'record_type': base.record_type, 'status': base.status}
for word in ['no trouble', 'cancel', 'duplicate', 'install', 'preventa', 'pm', 'contract']:
    for col, s in label_cols.items():
        hits = s.dropna()[s.dropna().str.lower().str.contains(r'\b' + word, regex=True)]
        if len(hits):
            facts[f'label_hit[{word}] in {col}'] = '; '.join(f'{k}={v}' for k, v in hits.value_counts().items())
for st_name, g in base.groupby('status'):
    facts[f'status[{st_name}] resolution_present / dispatch_no_present / n'] = (
        f"{int(g.resolution_present.sum())} / {int(raw.loc[g.index, 'Dispatch #'].map(lambda x: clean(x) is not None).sum())} / {len(g)}")
# Text-vs-label consistency (aggregate counts only; free text is never written).
notes = (raw['Trouble Resolution'].fillna('').astype(str) + ' ' + raw['Administrative Notes'].fillna('').astype(str)).str.lower()
NOT_REPRO = (r"no (?:trouble|problem|issue)s? found|\bntf\b|(?:could not|couldn't|unable to|not able to|never able to|"
             r"did not|didn't|cannot|can't|could't) (?:duplicate|replicate|reproduce)|problem not duplicated")
DUP_CASE = r"duplicate (?:case|call|work order|ticket|dispatch)|(?:closed|closing) (?:it )?as (?:a )?duplicate|because it(?:'|&#39;)?s duplicated"
for label, pattern, cat in [('NTF / could-not-reproduce wording', NOT_REPRO, 'No Trouble Found'),
                            ('cancel wording (upper bound; may cite another case)', r'\bcancel', 'Cancelled Request'),
                            ('duplicate-case wording', DUP_CASE, 'Duplicate Case')]:
    has = notes.str.contains(pattern, regex=True)
    facts[f'{label} in resolution/admin notes: tickets'] = int(has.sum())
    facts[f'{label} in resolution/admin notes: labelled {cat}'] = int((has & (base.category == cat)).sum())
    facts[f'{label} in resolution/admin notes: NOT labelled {cat}, in unplanned population'] = int(
        (has & (base.category != cat) & base.case_id.isin(service.case_id)).sum())
    facts[f'labelled {cat} without such wording'] = int((~has & (base.category == cat)).sum())
for k, v in facts.items():
    print(f'{k}: {v}')
out(pd.DataFrame(list(facts.items()), columns=['fact', 'value']), 'service_group_facts')

# Proactive monitoring tickets by service group / record type.
pm_tab = pd.crosstab(base.service_group, base.proactive_monitoring, margins=True).reset_index()
pm_tab.columns = ['service_group', 'customer_or_other', 'proactive_monitoring', 'all']
print(pm_tab.to_string(index=False))
print(pd.crosstab(base.record_type.fillna('[NULL]'), base.proactive_monitoring).to_string())
out(pm_tab, 'proactive_monitoring_by_group')

# --------------------------------------------------------------------------------------------
banner('2. Population sensitivity: unplanned n and repeat rate under re-classification')
# --------------------------------------------------------------------------------------------
variants = {
    'V0 Zac baseline (cell 15/31)': service,
    'V1 drop open (status Dispatch)': service[service.status != 'Dispatch'],
    'V2 drop all non-Trouble-Closed statuses': service[service.status == 'Trouble Closed'],
    'V3 drop No Trouble Found': service[service.service_group != 'No Trouble Found'],
    'V4 drop Missing trouble code': service[service.service_group != 'Missing trouble code'],
    'V5 drop proactive monitoring tickets': service[service.proactive_monitoring == 0],
    'V6 drop record type General Case': service[service.record_type != 'General Case'],
    'V7 = V2 + V3 + V5 (closed, coded-as-found, customer-initiated)':
        service[(service.status == 'Trouble Closed') & (service.service_group != 'No Trouble Found') &
                (service.proactive_monitoring == 0)],
    'V8 = V7 + drop Missing trouble code + General Case':
        service[(service.status == 'Trouble Closed') & (service.service_group == 'Other recorded issue') &
                (service.proactive_monitoring == 0) & (service.record_type != 'General Case')],
}
rows = []
for name, frame in variants.items():
    s = repeat_summary(frame)
    s['variant'] = name
    s['delta_repeat_pct_pts'] = round(s['repeat_pct'] - BASE['repeat_pct'], 2)
    rows.append(s)
sens = pd.DataFrame(rows)[['variant', 'unplanned_population', 'eligible_tickets', 'machines', 'repeat_tickets',
                           'repeat_pct', 'delta_repeat_pct_pts', 'same_category_repeats', 'same_category_repeat_pct']]
print(sens.to_string(index=False))

# --------------------------------------------------------------------------------------------
banner('3. Cell 31 repeat contact: gaps, unlabelled duplicates, concentration, censoring')
# --------------------------------------------------------------------------------------------
r = baseline_rep.copy()
r['opened'] = pd.to_datetime(r.opened_at)
r['next'] = pd.to_datetime(r.next_opened)
r['gap_min'] = (r['next'] - r['opened']).dt.total_seconds() / 60
# Exact-minute check of the floating julianday BETWEEN 0 AND 30.
exact = r.gap_min.between(0, 30 * 1440)
print('julianday vs exact-minute repeat flag disagreements:', int((exact.astype(int) != r.repeat_30).sum()))
print('gaps within 1 minute of exactly 30 days:', int(((r.gap_min - 43200).abs() <= 1).sum()))

bins = [-1, 0, 60, 240, 1440, 2 * 1440, 7 * 1440, 14 * 1440, 30 * 1440]
labels = ['0 (tie)', '(0,1h]', '(1h,4h]', '(4h,1d]', '(1d,2d]', '(2d,7d]', '(7d,14d]', '(14d,30d]']
rep = r[r.repeat_30 == 1].copy()
rep['gap_bin'] = pd.cut(rep.gap_min, bins=bins, labels=labels, right=True)
nxt_cat_missing = rep.next_category.isna()
gap_tab = rep.groupby('gap_bin', observed=False).agg(
    repeats=('case_id', 'size'), same_category=('same_category_30', 'sum')).reset_index()
gap_tab['cum_repeats'] = gap_tab.repeats.cumsum()
gap_tab['share_of_repeats_pct'] = (100 * gap_tab.repeats / len(rep)).round(2)
print(gap_tab.to_string(index=False))
out(gap_tab.astype({'gap_bin': str}), 'repeat_gap_bins')

# Same machine pairs that look like unlabelled duplicates.
dup_like = {}
for hours in (1, 4, 24):
    m = rep.gap_min <= hours * 60
    both_cat = rep.category.notna() & rep.next_category.notna()
    dup_like[f'repeats_gap_le_{hours}h'] = int(m.sum())
    dup_like[f'repeats_gap_le_{hours}h_same_category'] = int((m & (rep.category == rep.next_category)).sum())
    dup_like[f'repeats_gap_le_{hours}h_different_nonmissing_category'] = int(
        (m & both_cat & (rep.category != rep.next_category)).sum())
    dup_like[f'repeats_gap_le_{hours}h_one_side_NTF'] = int(
        (m & ((rep.category == 'No Trouble Found') | (rep.next_category == 'No Trouble Found'))).sum())
    dup_like[f'repeats_gap_le_{hours}h_one_side_missing_category'] = int(
        (m & (rep.category.isna() | rep.next_category.isna())).sum())
# Labelled duplicates for scale: where do the 15 'Duplicate Case' tickets sit relative to their predecessor?
allm = base.sort_values(['machine_id', 'opened_at', 'case_id']).copy()
allm['opened'] = pd.to_datetime(allm.opened_at)
allm['prev_gap_min'] = allm.groupby('machine_id').opened.diff().dt.total_seconds() / 60
dups = allm[allm.service_group == 'Duplicate case']
dup_like['labelled_duplicates'] = len(dups)
dup_like['labelled_duplicates_with_prior_ticket_same_machine'] = int(dups.prev_gap_min.notna().sum())
dup_like['labelled_duplicates_prior_gap_le_1d'] = int((dups.prev_gap_min <= 1440).sum())
dup_like['labelled_duplicates_prior_gap_median_hours'] = round(float(dups.prev_gap_min.median() / 60), 2)
allm['next_gap_min'] = -allm.groupby('machine_id').opened.diff(-1).dt.total_seconds() / 60
dup_like['labelled_duplicates_with_next_ticket_same_machine'] = int(dups.case_id.isin(allm[allm.next_gap_min.notna()].case_id).sum())
dup_like['labelled_duplicates_next_or_prior_gap_le_1d'] = int(
    ((allm.service_group == 'Duplicate case') & ((allm.prev_gap_min <= 1440) | (allm.next_gap_min <= 1440))).sum())
tie = allm[allm.prev_gap_min == 0]
dup_like['same_machine_same_timestamp_pairs_all_tickets'] = len(tie)
dup_like['same_timestamp_pair_groups'] = '; '.join(
    f"{allm.loc[allm.index[allm.index.get_loc(i) - 1], 'service_group']} + {allm.loc[i, 'service_group']}" for i in tie.index)
for k, v in dup_like.items():
    print(f'{k}: {v}')


def collapse(frame, hours):
    """Drop any service ticket opened within `hours` of the previous service ticket on the machine."""
    f = frame.sort_values(['machine_id', 'opened_at', 'case_id']).copy()
    t = pd.to_datetime(f.opened_at)
    gap = t.groupby(f.machine_id).diff().dt.total_seconds() / 3600
    return f[~(gap <= hours)]


for hours in (1, 24):
    s = repeat_summary(collapse(service, hours))
    s['variant'] = f'C{hours} collapse follow-ups within {hours}h of previous ticket (episode de-dup)'
    s['delta_repeat_pct_pts'] = round(s['repeat_pct'] - BASE['repeat_pct'], 2)
    sens = pd.concat([sens, pd.DataFrame([s])[sens.columns]], ignore_index=True)

# Same-category: next ticket only vs any later ticket within 30 days.
svc = service.sort_values(['machine_id', 'opened_at', 'case_id']).copy()
svc['opened'] = pd.to_datetime(svc.opened_at)
any_same, next_same = [], []
for mid, g in svc.groupby('machine_id', sort=False):
    t = g.opened.to_numpy()
    cat = [x if isinstance(x, str) else None for x in g.category]
    for i in range(len(g)):
        if g.opened_at.iloc[i] >= '2026-08-02':
            any_same.append(None)
            next_same.append(None)
            continue
        hit = 0
        nxt_hit = 0
        if cat[i] is not None:
            for j in range(i + 1, len(g)):
                if (t[j] - t[i]) > np.timedelta64(30, 'D'):
                    break
                if cat[j] == cat[i]:
                    hit = 1
                    nxt_hit = int(j == i + 1)
                    break
        any_same.append(hit)
        next_same.append(nxt_hit)
svc['same_cat_next_30'] = next_same
svc['same_cat_any_30'] = any_same
elig = svc[svc.opened_at < '2026-08-02']
same_any = int(elig.same_cat_any_30.sum())
assert int(elig.same_cat_next_30.sum()) == BASE['same_category_repeats'], 'loop does not reproduce next-only count'
print(f'Same-category repeat (any ticket within 30 d): {same_any} = {100 * same_any / len(elig):.2f}% '
      f'vs next-ticket-only {BASE["same_category_repeats"]} = {BASE["same_category_repeat_pct"]}%')
sens = pd.concat([sens, pd.DataFrame([{**BASE, 'variant': 'S1 same-category uses ANY ticket within 30 d (not just next)',
                                       'same_category_repeats': same_any,
                                       'same_category_repeat_pct': round(100 * same_any / len(elig), 2),
                                       'delta_repeat_pct_pts': 0.0}])[sens.columns]], ignore_index=True)
out(sens, 'population_sensitivity')

# Concentration of repeats across machines.
per_machine = r.groupby('machine_id').agg(eligible=('case_id', 'size'), repeats=('repeat_30', 'sum'))
pm_sorted = per_machine.sort_values(['repeats', 'eligible'], ascending=False)
n_m = len(pm_sorted)
conc = []
for share in (0.01, 0.05, 0.10, 0.20):
    k = math.ceil(share * n_m)
    top = pm_sorted.head(k)
    conc.append({'top_share_of_machines': f'{share:.0%}', 'machines': k, 'repeats': int(top.repeats.sum()),
                 'share_of_all_repeats_pct': round(100 * top.repeats.sum() / per_machine.repeats.sum(), 2),
                 'eligible_tickets': int(top.eligible.sum()),
                 'share_of_eligible_pct': round(100 * top.eligible.sum() / per_machine.eligible.sum(), 2)})
conc_df = pd.DataFrame(conc)
machines_with_repeat = int((per_machine.repeats > 0).sum())
extra = {
    'machines_with_eligible_ticket': n_m,
    'machines_with_ge1_repeat': machines_with_repeat,
    'machines_with_ge1_repeat_pct': round(100 * machines_with_repeat / n_m, 2),
    'machines_with_exactly_1_eligible_ticket': int((per_machine.eligible == 1).sum()),
    'machine_weighted_mean_repeat_pct': round(100 * (per_machine.repeats / per_machine.eligible).mean(), 2),
    'max_repeats_one_machine': int(per_machine.repeats.max()),
    'repeat_pct_excluding_top5pct_machines': round(
        100 * pm_sorted.iloc[math.ceil(.05 * n_m):].repeats.sum() / pm_sorted.iloc[math.ceil(.05 * n_m):].eligible.sum(), 2),
}
print(conc_df.to_string(index=False))
for k, v in extra.items():
    print(f'{k}: {v}')
conc_df = pd.concat([conc_df, pd.DataFrame([{'top_share_of_machines': k, 'machines': v} for k, v in extra.items()])],
                    ignore_index=True)
out(conc_df, 'repeat_concentration')


def wilson(k, n, z=1.96):
    p = k / n
    d = 1 + z * z / n
    c = (p + z * z / (2 * n)) / d
    h = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n)) / d
    return round(100 * (c - h), 1), round(100 * (c + h), 1)


mrows = []
for model, g in r.groupby('model'):
    pmm = g.groupby('machine_id').repeat_30.sum().sort_values(ascending=False)
    lo, hi = wilson(int(g.repeat_30.sum()), len(g))
    mrows.append({'model': model, 'eligible_tickets': len(g), 'repeat_tickets': int(g.repeat_30.sum()),
                  'repeat_pct': round(100 * g.repeat_30.mean(), 2), 'wilson95_lo': lo, 'wilson95_hi': hi,
                  'machines': g.machine_id.nunique(), 'repeats_top1_machine': int(pmm.iloc[0]),
                  'repeats_top3_machines': int(pmm.head(3).sum()),
                  'repeat_pct_without_top3_machines': round(
                      100 * g[~g.machine_id.isin(pmm.head(3).index)].repeat_30.mean(), 2)})
model_rep = pd.DataFrame(mrows).sort_values('eligible_tickets', ascending=False)
print(model_rep.head(12).to_string(index=False))
out(model_rep, 'repeat_by_model_robustness')

# Composition of the 1,574 repeats by what the index / follow-up ticket is.
nxt = svc[['machine_id', 'case_id', 'status', 'service_group', 'proactive_monitoring', 'record_type']].copy()
for col in ['case_id', 'status', 'service_group', 'proactive_monitoring', 'record_type']:
    nxt['next_' + col] = nxt.groupby('machine_id')[col].shift(-1)
rep_c = rep.merge(nxt[['case_id'] + [c for c in nxt.columns if c.startswith('next_')]], on='case_id', how='left')
comp = {
    'repeats': len(rep_c),
    'index_is_NTF': int((rep_c.service_group == 'No Trouble Found').sum()),
    'followup_is_NTF': int((rep_c.next_service_group == 'No Trouble Found').sum()),
    'either_is_NTF': int(((rep_c.service_group == 'No Trouble Found') | (rep_c.next_service_group == 'No Trouble Found')).sum()),
    'index_missing_code': int((rep_c.service_group == 'Missing trouble code').sum()),
    'followup_missing_code': int((rep_c.next_service_group == 'Missing trouble code').sum()),
    'followup_status_not_Trouble_Closed': int((rep_c.next_status != 'Trouble Closed').sum()),
    'followup_status_Dispatch': int((rep_c.next_status == 'Dispatch').sum()),
    'index_proactive_monitoring': int((rep_c.proactive_monitoring == 1).sum()),
    'followup_proactive_monitoring': int((rep_c.next_proactive_monitoring == 1).sum()),
    'either_proactive_monitoring': int(((rep_c.proactive_monitoring == 1) | (rep_c.next_proactive_monitoring == 1)).sum()),
    'followup_record_type_General_Case': int((rep_c.next_record_type == 'General Case').sum()),
    'followup_category_missing': int(nxt_cat_missing.sum()),
}
for k, v in comp.items():
    print(f'{k}: {v}')
out(pd.DataFrame(list(comp.items()), columns=['repeat_composition', 'repeats']), 'repeat_composition')

# Right-censoring / truncation rule.
EXPORT_LAST_OPENED = base.opened_at.max()
with zipfile.ZipFile(XLS) as z:
    core = z.read('docProps/core.xml').decode()
created = re.search(r'<dcterms:created[^>]*>([^<]+)<', core).group(1)
modified = re.search(r'<dcterms:modified[^>]*>([^<]+)<', core).group(1)
idx = r.copy()
idx['window_end'] = idx.opened + pd.Timedelta(days=30)
cens = {
    'export_last_opened_at': EXPORT_LAST_OPENED,
    'workbook_docprops_created_utc': created,
    'workbook_docprops_modified_utc': modified,
    'eligibility_rule': "opened_at < '2026-08-02' (string compare; includes all of 2026-08-01)",
    'eligible_index_opened_2026_08_01': int(idx.opened_at.str.startswith('2026-08-01').sum()),
    'eligible_index_window_ends_after_last_opened_ticket': int((idx.window_end > pd.Timestamp(EXPORT_LAST_OPENED)).sum()),
    'eligible_index_window_ends_after_2026_08_31_2359': int((idx.window_end > pd.Timestamp('2026-08-31 23:59:59')).sum()),
    'index_tickets_opened_2026_08_02_to_08_31_excluded': int(((service.opened_at >= '2026-08-02')).sum()),
}
# what if no censoring cut at all (index tickets up to 31 Aug) -> downward bias demonstration
allidx = sql(REPEAT_VIEW_SQL.replace("WHERE opened_at<'2026-08-02'", '').format(src='service_cases'))
cens['repeat_pct_without_censoring_cut'] = round(100 * allidx.repeat_30.mean(), 2)
cens['eligible_without_censoring_cut'] = len(allidx)
# Index tickets in the first 30 days of the export (left-truncated history) vs the rest.
early = r.opened_at < '2025-09-01'
cens['repeat_pct_index_opened_2025_08'] = round(100 * r[early].repeat_30.mean(), 2)
cens['repeat_pct_index_opened_2025_09_onward'] = round(100 * r[~early].repeat_30.mean(), 2)
for k, v in cens.items():
    print(f'{k}: {v}')
out(pd.DataFrame(list(cens.items()), columns=['item', 'value']), 'repeat_censoring')

# IPA's own Summary sheet names one 'most repeated serial'. Map it to our pseudonym only.
wbv = openpyxl.load_workbook(XLS, read_only=True, data_only=True)['Summary']
summary_rows = [row for row in wbv.iter_rows(values_only=True) if any(v is not None for v in row)]
ipa_serial = str(summary_rows[-1][0]).strip()
ipa_serial_cases = summary_rows[-1][1]
serial_counts = raw['Serial #'].map(clean).value_counts()
ipa_serial_pseudo = maps['serial'].get(ipa_serial)
ipa_machine = base[base.machine_id == ipa_serial_pseudo]
print(f'IPA Summary most-repeated serial -> {ipa_serial_pseudo}: Summary cases {ipa_serial_cases}, '
      f'our count {len(ipa_machine)}, max tickets on any serial {serial_counts.max()}, '
      f'serials tied at max {int((serial_counts == serial_counts.max()).sum())}; '
      f'its service groups {ipa_machine.service_group.value_counts().to_dict()}, '
      f'proactive monitoring tickets {int(ipa_machine.proactive_monitoring.sum())}, '
      f'repeats as index {int(r[r.machine_id == ipa_serial_pseudo].repeat_30.sum())}, '
      f'model {ipa_machine.model.unique().tolist()}')

# --------------------------------------------------------------------------------------------
banner('4. Cell 20 issue_mix vs IPA Summary sheet')
# --------------------------------------------------------------------------------------------
summary_top = [(row[1], row[2]) for row in summary_rows[2:12]]
full_counts = base.category.value_counts()
svc_counts = service.category.fillna('[Missing]').value_counts()
planned_counts = base[base.service_group == 'Planned work'].category.value_counts()
zac_rank = {cat: i + 1 for i, cat in enumerate(svc_counts.index)}
cmp_rows = []
for rank, (cat, n) in enumerate(summary_top, 1):
    cmp_rows.append({'category': cat, 'ipa_summary_rank': rank, 'ipa_summary_cases': n,
                     'full_export_count': int(full_counts.get(cat, 0)),
                     'zac_issue_mix_cases': int(svc_counts.get(cat, 0)), 'zac_issue_mix_rank': zac_rank.get(cat),
                     'diff_ipa_minus_zac': int(n - svc_counts.get(cat, 0)),
                     'planned_work_tickets_in_category': int(planned_counts.get(cat, 0))})
for cat in svc_counts.index[:12]:
    if cat not in dict(summary_top):
        cmp_rows.append({'category': cat, 'ipa_summary_rank': None, 'ipa_summary_cases': None,
                         'full_export_count': int(full_counts.get(cat, 0)) if cat != '[Missing]' else int(base.category.isna().sum()),
                         'zac_issue_mix_cases': int(svc_counts[cat]), 'zac_issue_mix_rank': zac_rank[cat],
                         'diff_ipa_minus_zac': None, 'planned_work_tickets_in_category': int(planned_counts.get(cat, 0))})
cmp = pd.DataFrame(cmp_rows)
print(cmp.to_string(index=False))
print('IPA top-10 sum (cached):', summary_rows[12][2], '| recomputed from full export:',
      int(sum(full_counts.get(c, 0) for c, _ in summary_top)))
print('Summary formulas are live COUNTIFs over the whole case sheet (all statuses / groups).')
out(cmp, 'issue_mix_vs_ipa_summary')

missing_comp = service[service.category.isna()].assign(
    trouble_code=lambda d: d.trouble_code.fillna('[Missing]'),
    trouble_closed=lambda d: (d.status == 'Trouble Closed').astype(int),
    proactive=lambda d: d.proactive_monitoring)
mc = missing_comp.groupby(['trouble_code', 'status']).size().reset_index(name='cases')
print(mc.to_string(index=False))
out(mc, 'missing_category_composition')
alt_den = len(service) - int((service.service_group == 'No Trouble Found').sum()) - int(service.category.isna().sum())
print(f'Deposit Door share: {100 * 1124 / len(service):.2f}% of 7,186 (Zac) vs '
      f'{100 * 1124 / alt_den:.2f}% of {alt_den} tickets with a recorded non-NTF category')

# --------------------------------------------------------------------------------------------
banner('5. Cell 13 field completeness: placeholders, structural missingness')
# --------------------------------------------------------------------------------------------
PLACEHOLDERS = {'.', '-', '--', '...', 'n/a', 'na', 'n.a.', 'none', 'null', 'tbd', '?', 'x', '0', 'unknown'}
ph_rows = []
for col in raw.columns:
    v = raw[col].map(lambda x: None if x is None else str(x).strip())
    present = v.notna() & (v != '')
    ph = present & v.str.lower().isin(PLACEHOLDERS)
    ph_rows.append({'column': col, 'present_after_clean': int(present.sum()),
                    'placeholder_values_counted_present': int(ph.sum()),
                    'whitespace_only_cells': int((v == '').sum()),
                    'values_with_outer_whitespace': int(sum(1 for x in raw0[col] if isinstance(x, str) and x != x.strip()))})
ph_df = pd.DataFrame(ph_rows)
print(ph_df[ph_df[['placeholder_values_counted_present', 'whitespace_only_cells', 'values_with_outer_whitespace']].sum(axis=1) > 0]
      .to_string(index=False))
out(ph_df, 'placeholder_values')
text_len = {
    'trouble_reported_len_lt_10': int((cases.reported_length < 10).sum()),
    'trouble_reported_len_lt_20': int((cases.reported_length < 20).sum()),
    'trouble_resolution_present_len_lt_10': int(((cases.resolution_present == 1) & (cases.resolution_length < 10)).sum()),
    'trouble_reported_median_len': float(cases.reported_length.median()),
}
hw = base.trouble_code.isin(['Electro/Mechanical', 'Electronics'])
text_len['subcategory_missing_E/M+Electronics'] = f"{int(base[hw].subcategory.isna().sum())}/{int(hw.sum())} = {100 * base[hw].subcategory.isna().mean():.2f}%"
text_len['subcategory_missing_other_codes_incl_missing'] = f"{int(base[~hw].subcategory.isna().sum())}/{int((~hw).sum())} = {100 * base[~hw].subcategory.isna().mean():.2f}%"
text_len['trouble_code_missing_among_Trouble_Closed'] = f"{int(base[base.status == 'Trouble Closed'].trouble_code.isna().sum())}/{int((base.status == 'Trouble Closed').sum())}"
text_len['trouble_closed_missing_code_proactive_monitoring'] = int(((base.status == 'Trouble Closed') & base.trouble_code.isna() & (base.proactive_monitoring == 1)).sum())
text_len['trouble_closed_missing_code_General_Case'] = int(((base.status == 'Trouble Closed') & base.trouble_code.isna() & (base.record_type == 'General Case')).sum())
for k, v in text_len.items():
    print(f'{k}: {v}')
out(pd.DataFrame(list(text_len.items()), columns=['item', 'value']), 'text_and_hierarchy_facts')

# Missingness by status and by service group: how much of the 101/330/2,315/641 is structural or WIP.
fc_rows = []
for label, mask in [('All tickets', base.case_id.notna()),
                    ('Status Trouble Closed', base.status == 'Trouble Closed'),
                    ('Status not Trouble Closed', base.status != 'Trouble Closed'),
                    ('Unplanned (service_cases)', base.case_id.isin(service.case_id)),
                    ('Other recorded issue & Trouble Closed', (base.service_group == 'Other recorded issue') & (base.status == 'Trouble Closed')),
                    ('No Trouble Found', base.service_group == 'No Trouble Found'),
                    ('Cancelled request', base.service_group == 'Cancelled request'),
                    ('Proactive monitoring', base.proactive_monitoring == 1)]:
    b = base[mask]
    fc_rows.append({'subset': label, 'tickets': len(b),
                    'trouble_code_missing': int(b.trouble_code.isna().sum()),
                    'category_missing': int(b.category.isna().sum()),
                    'subcategory_missing': int(b.subcategory.isna().sum()),
                    'subcategory_missing_pct': round(100 * b.subcategory.isna().mean(), 2),
                    'resolution_missing': int((b.resolution_present == 0).sum()),
                    'resolution_missing_pct': round(100 * (b.resolution_present == 0).mean(), 2)})
fc = pd.DataFrame(fc_rows)
print(fc.to_string(index=False))
out(fc, 'field_completeness_by_subset')
sub_by_code = base.assign(sub_missing=base.subcategory.isna().astype(int)).groupby(
    base.trouble_code.fillna('[Missing]')).agg(tickets=('case_id', 'size'), subcategory_missing=('sub_missing', 'sum'))
sub_by_code['subcategory_missing_pct'] = (100 * sub_by_code.subcategory_missing / sub_by_code.tickets).round(1)
sub_by_rt = base.assign(sub_missing=base.subcategory.isna().astype(int)).groupby(
    base.record_type.fillna('[Missing]')).agg(tickets=('case_id', 'size'), subcategory_missing=('sub_missing', 'sum'))
sub_by_rt['subcategory_missing_pct'] = (100 * sub_by_rt.subcategory_missing / sub_by_rt.tickets).round(1)
sub_by_cat = base[base.category.notna()].assign(sub_missing=lambda d: d.subcategory.isna().astype(int)).groupby('category').agg(
    tickets=('case_id', 'size'), subcategory_missing=('sub_missing', 'sum'))
sub_by_cat['subcategory_missing_pct'] = (100 * sub_by_cat.subcategory_missing / sub_by_cat.tickets).round(1)
print(sub_by_code.sort_values('tickets', ascending=False).to_string())
print(sub_by_rt.sort_values('tickets', ascending=False).to_string())
print(sub_by_cat.sort_values('subcategory_missing', ascending=False).head(15).to_string())
sub_tables = pd.concat([sub_by_code.reset_index().rename(columns={'trouble_code': 'value'}).assign(dimension='trouble_code'),
                        sub_by_rt.reset_index().rename(columns={'record_type': 'value'}).assign(dimension='record_type'),
                        sub_by_cat.reset_index().rename(columns={'category': 'value'}).assign(dimension='category')])
out(sub_tables[['dimension', 'value', 'tickets', 'subcategory_missing', 'subcategory_missing_pct']], 'subcategory_missing_by')

# Hierarchy: is category -> trouble_code a strict parent? (ambiguity only arises at the subcategory level)
cat_parents = base.dropna(subset=['category', 'trouble_code']).groupby('category').trouble_code.nunique()
print('categories with >1 parent trouble code:', int((cat_parents > 1).sum()), 'of', len(cat_parents))
amb = sql("""SELECT trouble_code, subcategory, COUNT(DISTINCT category) categories, COUNT(*) cases
FROM cases WHERE trouble_code IS NOT NULL AND subcategory IS NOT NULL AND category IS NOT NULL
GROUP BY trouble_code,subcategory HAVING COUNT(DISTINCT category)>1 ORDER BY cases DESC""")
print('hierarchy ambiguity rows:', len(amb), amb.cases.tolist())
amb_all = base.dropna(subset=['subcategory', 'category']).groupby('subcategory').category.nunique()
print('subcategories (ignoring trouble code) spanning >1 category:', int((amb_all > 1).sum()), 'of', len(amb_all))

# --------------------------------------------------------------------------------------------
banner('6. Dates (cell 4 conversion)')
# --------------------------------------------------------------------------------------------
close_raw = raw['Trouble Closed Date/Time']
open_raw = raw['Date/Time Opened']
zac_close = pd.to_datetime(close_raw).dt.strftime('%Y-%m-%d %H:%M:%S').where(close_raw.notna(), None)
zac_open = pd.to_datetime(open_raw).dt.strftime('%Y-%m-%d %H:%M:%S').where(open_raw.notna(), None)
op = pd.to_datetime(cases.opened_at)
cl = pd.to_datetime(cases.closed_at)
dur_h = (cl - op).dt.total_seconds() / 3600
dates = {
    'opened_values_by_python_type': str(open_raw.map(lambda v: type(v).__name__).value_counts().to_dict()),
    'closed_values_by_python_type': str(close_raw.map(lambda v: type(v).__name__).value_counts().to_dict()),
    'closed_text_values': int(close_raw.map(lambda v: isinstance(v, str)).sum()),
    'closed_text_value_parses_to': cases.closed_at[close_raw.map(lambda v: isinstance(v, str))].iloc[0],
    'zac_vectorised_equals_elementwise_opened': bool((zac_open.fillna('~') == cases.opened_at.fillna('~')).all()),
    'zac_vectorised_equals_elementwise_closed': bool((zac_close.fillna('~') == cases.closed_at.fillna('~')).all()),
    'opened_nonzero_seconds_or_microseconds': int(sum(1 for v in open_raw if v.second or v.microsecond)),
    'closed_nonzero_seconds_or_microseconds': int(sum(1 for v in close_raw if hasattr(v, 'second') and (v.second or v.microsecond))),
    'closed_missing': int(cases.closed_at.isna().sum()),
    'closed_missing_status_breakdown': str(cases[cases.closed_at.isna()].status.value_counts().to_dict()),
    'closed_before_opened': int((dur_h < 0).sum()),
    'closed_equal_opened': int((dur_h == 0).sum()),
    'closed_within_1h': int(((dur_h >= 0) & (dur_h < 1)).sum()),
    'median_open_to_close_hours': round(float(dur_h.median()), 2),
    'closed_after_last_opened_ticket': int((cl > op.max()).sum()),
    'latest_close': str(cl.max()),
    'closed_at_used_in_any_notebook_cell_after_4': False,
    'status_used_in_any_notebook_cell_after_4': False,
}
nb = (Path(__file__).resolve().parents[1] / 'data' / 'derived' / 'notebook_dump.md').read_text()
after4 = nb.split('## Cell 5 [markdown]')[1]
code_after4 = '\n'.join(re.findall(r'```python\n(.*?)```', after4, flags=re.S))
dates['closed_at_used_in_any_notebook_cell_after_4'] = 'closed_at' in code_after4
dates['status_used_in_any_notebook_cell_after_4'] = bool(re.search(r"(?<![\w'])status\b", code_after4.replace(
    "'Operational status'", '')))
for k, v in dates.items():
    print(f'{k}: {v}')
out(pd.DataFrame(list(dates.items()), columns=['check', 'value']), 'dates')

# Opening hour by hospital time zone: are case timestamps local or a single export zone?
TZ = {**{s: 'Eastern' for s in 'NY GA NJ NC MA PA ON OH VA FL CT MD SC IN WV MI DC DE ME QC NH KY'.split()},
      **{s: 'Central' for s in 'TX IL WI LA MO IA NE MN ND MS AL AR OK KS TN'.split()},
      **{s: 'Mountain' for s in 'NM CO UT MT ID'.split()}, 'AZ': 'Arizona (no DST)',
      **{s: 'Pacific' for s in 'CA OR NV WA'.split()}, 'HI': 'Hawaii'}
hz = cases.assign(tz=cases.state.map(TZ).fillna('Unparsed'), hour=op.dt.hour + op.dt.minute / 60)
tz_tab = hz.groupby('tz').agg(tickets=('case_id', 'size'), median_open_hour=('hour', 'median'),
                              p10_open_hour=('hour', lambda s: s.quantile(.1)),
                              opened_before_07=('hour', lambda s: int((s < 7).sum())),
                              opened_after_19=('hour', lambda s: int((s >= 19).sum()))).round(2).reset_index()
print(tz_tab.to_string(index=False))
out(tz_tab, 'open_hour_by_timezone')

# --------------------------------------------------------------------------------------------
banner('7. Cells 17/18 customer workload: name variants and multi-account serials')
# --------------------------------------------------------------------------------------------
acc = raw['Account Name'].map(clean)
accounts = pd.Series(sorted(acc.unique()))
STOP = {'inc', 'llc', 'the', 'of', 'and', 'co', 'corp', 'corporation', 'ltd'}
GENERIC = STOP | {'hospital', 'medical', 'center', 'centre', 'health', 'healthcare', 'system', 'regional', 'campus',
                  'university', 'clinic', 'clinics', 'st', 'saint'}


def n1(s):
    return ' '.join(re.sub(r'[^a-z0-9 ]', ' ', s.lower().replace('&', ' and ')).split())


def n2(s):
    return ' '.join(w for w in n1(s).split() if w not in STOP)


def n3(s):
    return ' '.join(w for w in n1(s).split() if w not in GENERIC)


acct_rows = []
for name, f in [('case/punctuation/& folded', n1), ('+ drop Inc/LLC/The/of/and', n2),
                ('+ drop Hospital/Medical/Center/Health/University/Saint etc.', n3)]:
    keys = accounts.map(f)
    clusters = accounts.groupby(keys).size()
    multi = clusters[clusters > 1]
    acct_rows.append({'normalisation': name, 'account_labels': len(accounts), 'distinct_after': keys.nunique(),
                      'clusters_with_gt1_label': len(multi), 'labels_in_clusters': int(multi.sum())})
# Do aggressive (n3) collisions share a location? Same normalised name in different cities = different facilities.
loc_by_acc = pd.DataFrame({'a': acc, 'l': location}).drop_duplicates().set_index('a').l
keys3 = accounts.map(n3)
coll = accounts.groupby(keys3).apply(list)
coll = coll[coll.map(len) > 1]
same_loc_clusters = sum(1 for v in coll if len({loc_by_acc[a] for a in v}) == 1)
acct_rows.append({'normalisation': 'aggressive clusters whose labels share one City/State/Zip',
                  'clusters_with_gt1_label': same_loc_clusters})
# Fuzzy within-state similarity
st_by_acc = pd.DataFrame({'a': acc, 'st': cases.state}).drop_duplicates('a').set_index('a').st
pairs = 0
pairs_same_loc = 0
for st, names in accounts.groupby(accounts.map(st_by_acc)):
    names = list(names)
    for i in range(len(names)):
        for j in range(i + 1, len(names)):
            if difflib.SequenceMatcher(None, n1(names[i]), n1(names[j])).ratio() >= 0.85:
                pairs += 1
                pairs_same_loc += int(loc_by_acc[names[i]] == loc_by_acc[names[j]])
acct_rows.append({'normalisation': 'within-state name pairs with difflib ratio >= 0.85', 'clusters_with_gt1_label': pairs})
acct_rows.append({'normalisation': '... of which share the same City/State/Zip', 'clusters_with_gt1_label': pairs_same_loc})
acct_rows.append({'normalisation': 'locations (City/State/Zip) with >1 account label',
                  'clusters_with_gt1_label': int(pd.DataFrame({'a': acc, 'l': location}).groupby('l').a.nunique().gt(1).sum())})
acct_rows.append({'normalisation': 'account labels with >1 location',
                  'clusters_with_gt1_label': int(pd.DataFrame({'a': acc, 'l': location}).groupby('a').l.nunique().gt(1).sum())})
acct_rows.append({'normalisation': "account labels containing a mis-encoded '?' apostrophe",
                  'clusters_with_gt1_label': int(accounts.str.contains(r'\w\?\w', regex=True).sum())})
acct_df = pd.DataFrame(acct_rows)
print(acct_df.to_string(index=False))
out(acct_df, 'account_name_variants')

# Serials under more than one account label.
sa = cases[['machine_id', 'customer', 'opened_at', 'state']].copy()
multi_serials = sa.groupby('machine_id').customer.nunique().loc[lambda s: s > 1].index
ms_rows = []
for mid in multi_serials:
    g = sa[sa.machine_id == mid].sort_values('opened_at')
    seq = g.customer.tolist()
    switches = sum(1 for a, b in zip(seq, seq[1:]) if a != b)
    for cust, gg in g.groupby('customer', sort=False):
        ms_rows.append({'machine_id': mid, 'customer': cust, 'tickets': len(gg), 'first_opened': gg.opened_at.min()[:10],
                        'last_opened': gg.opened_at.max()[:10], 'state': gg.state.iloc[0],
                        'account_switches_in_time_order': switches})
ms = pd.DataFrame(ms_rows)
print(ms.to_string(index=False))
out(ms, 'multi_account_serials')
cust = sql("""SELECT customer,COUNT(*) cases,COUNT(DISTINCT machine_id) affected_machines FROM cases
GROUP BY customer ORDER BY cases DESC,customer""")
top10 = cust.head(10).customer.tolist()
print('multi-account serials touching a top-10 customer:', int(ms.customer.isin(top10).sum()),
      '| tickets on multi-account serials:', int(ms.tickets.sum()),
      '| customer-machine pairs added by double counting:', int(len(ms) - len(multi_serials)))
# Sensitivity: attribute each multi-account serial to its most recent account (a move) and recount top-10.
latest = sa.sort_values('opened_at').groupby('machine_id').customer.last()
sa2 = sa.assign(customer=np.where(sa.machine_id.isin(multi_serials), sa.machine_id.map(latest), sa.customer))
cust2 = sa2.groupby('customer').agg(cases=('machine_id', 'size'), affected_machines=('machine_id', 'nunique')) \
    .reset_index().sort_values(['cases', 'customer'], ascending=[False, True])
print('top-10 unchanged after re-attributing multi-account serials:', cust2.head(10).customer.tolist() == top10,
      '| top-10 share', round(100 * cust2.head(10).cases.sum() / len(sa2), 2))
# System-level heuristic: label text before ' - ' (e.g. 'System - Campus'). Only a rough upper bound.
sys_key = acc.map(lambda s: s.split(' - ')[0].strip() if ' - ' in s else s)
sys_counts = sys_key.value_counts()
print(f"labels with a ' - ' system prefix: {int(accounts.str.contains(' - ').sum())}; prefix groups formed: "
      f"{int((accounts[accounts.str.contains(' - ')].map(lambda s: s.split(' - ')[0].strip()).value_counts() > 1).sum())}; "
      f"top-10 share if grouped by prefix: {100 * sys_counts.head(10).sum() / len(acc):.2f}% (vs 11.16%)")

# --------------------------------------------------------------------------------------------
banner('8. Cell 8 pilot sampling')
# --------------------------------------------------------------------------------------------
PILOT = ['T00705', 'T00783', 'T00951', 'T00957', 'T01157', 'T01287', 'T01629', 'T01795', 'T01796', 'T01797',
         'T02128', 'T02229', 'T02268', 'T03075', 'T03079', 'T03326', 'T03959', 'T04699', 'T04907', 'T05185',
         'T05506', 'T05683', 'T05924', 'T06107', 'T06354', 'T06445', 'T06663', 'T06672', 'T06712', 'T06912']
NO_TX = {'T01629', 'T03959', 'T05506', 'T05683'}  # 'No transactions observed' in Zac's cell-34 table
p = base[base.case_id.isin(PILOT)].copy()
assert len(p) == 30
p['stratum'] = p.trouble_code.fillna('Missing')
p['zac_availability'] = np.where(p.case_id.isin(NO_TX), 'No transactions observed', 'Candidate codes available')
p['repeat_index_any_followup_30d'] = p.case_id.map(baseline_rep.set_index('case_id').repeat_30)
pilot_out = p[['case_id', 'stratum', 'category', 'service_group', 'status', 'record_type', 'model',
               'proactive_monitoring', 'resolution_present', 'zac_availability']].assign(month=p.opened_at.str[:7]) \
    .sort_values(['stratum', 'case_id'])
print(pilot_out.to_string(index=False))
out(pilot_out, 'pilot_tickets')
win = base[base.opened_at.between('2025-09-02', '2026-07-31 23:59:59')].copy()
win['stratum'] = win.trouble_code.fillna('Missing')
strata = win.groupby('stratum').agg(window_tickets=('case_id', 'size')).reset_index()
strata['window_share_pct'] = (100 * strata.window_tickets / len(win)).round(2)
strata = strata.merge(p.groupby('stratum').size().rename('pilot_tickets').reset_index(), how='left').fillna({'pilot_tickets': 0})
strata['pilot_share_pct'] = (100 * strata.pilot_tickets / 30).round(2)
strata['over_representation_x'] = (strata.pilot_share_pct / strata.window_share_pct).round(1)
strata['window_planned_work'] = strata.stratum.map(win[win.service_group == 'Planned work'].groupby('stratum').size()).fillna(0).astype(int)
strata['window_not_unplanned'] = strata.stratum.map(win[~win.case_id.isin(service.case_id)].groupby('stratum').size()).fillna(0).astype(int)
print(strata.sort_values('window_tickets', ascending=False).to_string(index=False))
out(strata, 'pilot_strata_representation')
pil = {
    'pilot_planned_work': int((p.service_group == 'Planned work').sum()),
    'pilot_not_in_unplanned_population': int((~p.case_id.isin(service.case_id)).sum()),
    'pilot_service_groups': str(p.service_group.value_counts().to_dict()),
    'pilot_no_tx_service_groups': str(p[p.case_id.isin(NO_TX)].service_group.value_counts().to_dict()),
    'pilot_no_tx_strata': str(p[p.case_id.isin(NO_TX)].stratum.value_counts().to_dict()),
    'pilot_no_tx_record_types': str(p[p.case_id.isin(NO_TX)].record_type.value_counts().to_dict()),
    'pilot_proactive_monitoring': int(p.proactive_monitoring.sum()),
    'pilot_customer_initiated_recorded_issue': int(((p.service_group == 'Other recorded issue') & (p.proactive_monitoring == 0)).sum()),
    'pilot_no_tx_proactive_monitoring': int(p[p.case_id.isin(NO_TX)].proactive_monitoring.sum()),
    'pilot_record_type_General_Case': int((p.record_type == 'General Case').sum()),
    'population_proactive_monitoring_pct': round(100 * base.proactive_monitoring.mean(), 2),
    'window_proactive_monitoring_pct': round(100 * win.proactive_monitoring.mean(), 2),
    'pilot_status': str(p.status.value_counts().to_dict()),
    'pilot_resolution_missing': int((p.resolution_present == 0).sum()),
    'window_tickets': len(win),
    'window_strata': int(win.stratum.nunique()),
    'window_share_of_tickets_in_strata_E/M+Electronics': round(100 * win.stratum.isin(['Electro/Mechanical', 'Electronics']).mean(), 2),
    'pilot_share_E/M+Electronics': round(100 * p.stratum.isin(['Electro/Mechanical', 'Electronics']).mean(), 2),
    'window_share_rare_strata_PM_Install_RI_Canopy': round(100 * win.stratum.isin(['Preventative Maintenance', 'Installation', 'Repair/Investigate', 'Canopy']).mean(), 2),
    'pilot_share_rare_strata_PM_Install_RI_Canopy': round(100 * p.stratum.isin(['Preventative Maintenance', 'Installation', 'Repair/Investigate', 'Canopy']).mean(), 2),
    'lower_bound_reason': "2025-09-02 minus 32-day extraction buffer = 2025-08-01 (first telemetry date)",
}
lb = pd.Timestamp('2025-09-02') - pd.Timedelta(days=32)
pil['lower_bound_minus_32d'] = str(lb.date())
# String-window check: anchors have zero seconds/microseconds so isoformat(' ') is 'YYYY-MM-DD HH:MM:SS'.
import datetime
iso_ok = all(datetime.datetime.fromisoformat(t).isoformat(' ') == t for t in p.opened_at)
shift_ok = all(len((datetime.datetime.fromisoformat(t) + datetime.timedelta(hours=h, days=-30)).isoformat(' ')) == 19
               for t in p.opened_at for h in (-12, 0, 12))
pil['anchor_isoformat_roundtrip_equal'] = iso_ok
pil['shifted_window_bounds_all_19_chars'] = shift_ok
# Hash ordering independence check: pilot uses the pseudonym (T#####), so it changes if the case map changes.
pil['sample_key'] = "sha256('15072' + T#####) - depends on pseudonym order of Case Numbers"
for k, v in pil.items():
    print(f'{k}: {v}')
out(pd.DataFrame(list(pil.items()), columns=['item', 'value']), 'pilot_facts')

# --------------------------------------------------------------------------------------------
banner('9. Workbook checksum')
# --------------------------------------------------------------------------------------------
wb_rows = [{'item': 'sha256 here', 'value': sha256_file(XLS)}, {'item': 'sha256 Zac (cell 2)', 'value': ZAC_WORKBOOK_SHA256},
           {'item': 'docProps created (UTC)', 'value': created}, {'item': 'docProps modified (UTC)', 'value': modified},
           {'item': 'sheets', 'value': '; '.join(openpyxl.load_workbook(XLS, read_only=True).sheetnames)},
           {'item': 'Excel-derived printed numbers reproduced (this script, section 0)', 'value': bool(repro.match.all())}]
print(pd.DataFrame(wb_rows).to_string(index=False))
out(pd.DataFrame(wb_rows), 'workbook')
out(repro, 'baseline_reproduction')
print('\nDone.')
