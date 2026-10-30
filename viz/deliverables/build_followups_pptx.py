#!/usr/bin/env python3
"""The follow-up PowerPoint with real customer names: deliverables/IPA_Followups_30Oct2026_REALNAMES.pptx.

viz/deliverables/followups_pptx.js builds the follow-up deck from viz/deliverables/followups_pptx.json, and its
pseudonymous build names no customer. This real-name build adds the content's real_names_only section, the update
deck's customer slides 6 (sl_05 + sl_06) and 7 (sl_07 + sl_08), and, in order, failing loudly with pseudonyms and
counts only:
  1. every 'Customer NNN' in the titles, takeaways, card headers, alt text and speaker notes becomes its real name
     (names.substitute); a pseudonym without a name fails the build, and none may be left;
  2. each chart whose slide script draws a real-name variant (its REAL_PNG: sl_06 and sl_07) is the variant in
     deliverables/slides/, which must exist and differ from the pseudonymous chart;
  3. every takeaway is fitted to its bar's two lines: measured with Liberation Sans (Arial's metrics) at the bar's
     width, and set 0.5 pt smaller at a time where it needs more, down to MIN_PT;
then followups_pptx.js writes the deck from a temporary copy of that content (under deliverables/, gitignored, and
deleted afterwards, also on failure).

Checks on the written deck (counts only; any failure leaves no deck behind):
  - the slide count: cover, agenda, every content slide and the close;
  - no pseudonym in any slide text, speaker note or picture alt text, and the names found = the names substituted;
  - each swapped chart is embedded byte for byte (sha256 in ppt/media) and its pseudonymous chart is not;
  - the pptx skill's validator passes, when it is installed (only its summary line is printed).

Twin build (--twin OUT_DIR): the same deck with twin names (names.twin: each placeholder has its real name's word
lengths) and twin customer charts drawn by the slide scripts' plot_named(), into OUT_DIR, for checking the layout
without looking at real names.

Inputs:  viz/deliverables/followups_pptx.json; viz/figures/slides/; deliverables/slides/*_REALNAMES.png (real build);
         data/private/fx_private_customer_names.csv via names.real_names(); --assets DIR as for followups_pptx.js
Output:  deliverables/IPA_Followups_30Oct2026_REALNAMES.pptx (never commit), or OUT_DIR/..._TWIN.pptx
Requires Node with pptxgenjs (NODE_PATH), python-pptx and Pillow.
Run:  cd /home/user/analytics-lab && NODE_PATH=<node_modules with pptxgenjs> \\
      python3 -I viz/deliverables/build_followups_pptx.py --assets DIR [--twin OUT_DIR] [--tmp DIR]
"""
import argparse
import copy
import hashlib
import importlib
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path[:0] = [str(HERE), str(ROOT / 'viz' / 'slides')]  # python3 -I leaves the script directory off sys.path
import names  # noqa: E402

CONTENT = HERE / 'followups_pptx.json'
BUILDER = HERE / 'followups_pptx.js'
FIG = ROOT / 'viz' / 'figures' / 'slides'
OUT_REAL = names.DELIVERABLES / 'IPA_Followups_30Oct2026_REALNAMES.pptx'
TWIN_FILE = 'IPA_Followups_30Oct2026_TWIN.pptx'
FONTS = ['/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
         '/System/Library/Fonts/Supplemental/Arial.ttf', '/Library/Fonts/Arial.ttf', 'C:/Windows/Fonts/arial.ttf']
TAKEAWAY_PX, TAKEAWAY_PT, MIN_PT = 1000, 13.0, 11.0   # the takeaway box: 1000 px of the 1920-px design (1 px = 0.5 pt)
SLACK = 0.97                                          # wrap at 97% of the width: renderers break lines a little earlier
COVER_NOTE = 'Customer names do not appear in this deck.'
COVER_NOTE_REAL = ("This version adds the update deck's customer slides 6 and 7 with real customer names: for the "
                   'project team and IPA only.')


class Failed(Exception):
    """A step or check failed; the message names pseudonyms, counts, slide labels and paths only."""


def fail(message):
    raise Failed(message)


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def redactor(*maps):
    """Replace any name of the given maps in a message by '<name>' (longest first)."""
    found = sorted({str(v) for m in maps for v in m.values()}, key=len, reverse=True)

    def redact(text):
        for v in found:
            text = text.replace(v, '<name>')
        return text
    return redact


def substitute_all(obj, names_map, counter):
    """Every string in obj with its pseudonyms replaced by names_map's names (counter['n'] counts them)."""
    if isinstance(obj, dict):
        return {k: substitute_all(v, names_map, counter) for k, v in obj.items()}
    if isinstance(obj, list):
        return [substitute_all(v, names_map, counter) for v in obj]
    if isinstance(obj, str):
        text, n = names.substitute(obj, names_map)
        counter['n'] += n
        return text
    return obj


def strings(obj):
    if isinstance(obj, dict):
        return [s for v in obj.values() for s in strings(v)]
    if isinstance(obj, list):
        return [s for v in obj for s in strings(v)]
    return [obj] if isinstance(obj, str) else []


def lines_needed(text, pt, font_file):
    """Lines the text takes in the takeaway box at pt: greedy word wrap measured with the font, at 4x precision."""
    from PIL import ImageFont
    font = ImageFont.truetype(font_file, size=round(pt * 2 * 4))
    width, lines, current = TAKEAWAY_PX * 4 * SLACK, 1, ''
    for word in text.split():
        trial = f'{current} {word}'.strip()
        if font.getlength(trial) <= width:
            current = trial
        else:
            lines, current = lines + 1, word
    return lines


def fit_takeaways(slides, font_file):
    """Set takeaway_pt on each slide whose takeaway needs more than two lines at TAKEAWAY_PT. Returns the changes."""
    changed = []
    for s in slides:
        pt = TAKEAWAY_PT
        while lines_needed(s['takeaway'], pt, font_file) > 2 and pt > MIN_PT:
            pt -= 0.5
        if lines_needed(s['takeaway'], pt, font_file) > 2:
            fail(f"slide {s['label']}: the takeaway needs more than two lines even at {MIN_PT} pt")
        if pt < TAKEAWAY_PT:
            s['takeaway_pt'] = pt
            changed.append(f"{s['label']} {TAKEAWAY_PT:g}->{pt:g} pt")
    return changed


def customer_chart(png):
    """The slide-chart module drawing png, when it draws a real-name variant (defines REAL_PNG); else None."""
    stem = Path(png).stem
    if not (ROOT / 'viz' / 'slides' / f'{stem}.py').exists():
        return None
    module = importlib.import_module(stem)
    return module if hasattr(module, 'REAL_PNG') else None


def find_validator():
    if os.environ.get('PPTX_VALIDATE'):
        path = Path(os.environ['PPTX_VALIDATE'])
    else:
        found = sorted((Path.home() / '.claude' / 'skills').glob('**/pptx/scripts/office/validate.py'))
        path = found[0] if found else None
    return path if path is not None and path.is_file() else None


def build(assets, twin_dir=None, tmp_parent=None):
    real = twin_dir is None
    real_names = names.real_names()
    if real_names is None:
        fail(f'{names.NAMES_FILE.relative_to(names.ROOT)} not found: no names for the deck')
    names_map = real_names if real else names.twin(real_names)
    redact = redactor(real_names, names.twin(real_names))
    font_file = next((f for f in FONTS if Path(f).exists()), None)
    if font_file is None:
        fail('no Liberation Sans or Arial font file found to measure the takeaways')

    deck = json.loads(CONTENT.read_text())
    deck['real_names'] = True
    deck['file_title'] += ' (real customer names)' if real else ' (twin names)'
    if deck['cover']['notes'].count(COVER_NOTE) != 1:
        fail("the cover notes no longer say that customer names do not appear: update COVER_NOTE")
    deck['cover']['notes'] = deck['cover']['notes'].replace(COVER_NOTE, COVER_NOTE_REAL)
    extra = [sec for sec in deck['sections'] if sec.get('real_names_only')]
    if not extra:
        fail('the content has no real_names_only section')
    deck['agenda_notes'] += ' ' + ' '.join(f"Part {deck['sections'].index(sec) + 1}: {sec['label']}." for sec in extra)
    print(f"IPA follow-up PowerPoint ({'real names' if real else 'twin names'}): "
          f"{sum(len(sec['slides']) for sec in deck['sections'])} content slides in {len(deck['sections'])} sections, "
          f"{len(extra)} real-name-only")

    # 1. Names
    customers = sorted({c for s in strings(deck) for c in names.found(s)})
    counter = {'n': 0}
    try:
        deck = substitute_all(deck, names_map, counter)
    except KeyError as err:  # names pseudonyms only
        fail(f'step 1: {err.args[0]}')
    left = sum(names.pseudonyms_left(s) for s in strings(deck))
    if left:
        fail(f'step 1: {left} pseudonym(s) left after substitution')
    print(f"  step 1: {counter['n']} pseudonyms -> names ({len(customers)} customers: {', '.join(customers)}); "
          '0 pseudonyms left')

    # 2. Customer charts
    swaps = []
    for sec in deck['sections']:
        for s in sec['slides']:
            for c in s['cards']:
                module = customer_chart(c['png']) if c['kind'] == 'chart' else None
                if module is None:
                    continue
                stem = Path(c['png']).stem
                if real:
                    path = names.DELIVERABLES / 'slides' / module.REAL_PNG
                    if not path.is_file():
                        fail(f'step 2: {path.relative_to(ROOT)} not found: draw it with viz/slides/run_all.py')
                    if sha256(path) == sha256(FIG / c['png']):
                        fail(f'step 2: {path.name} is the same file as the pseudonymous {c["png"]}')
                else:
                    path = Path(twin_dir) / 'twin_charts' / f'{stem}_TWIN.png'
                    path.parent.mkdir(parents=True, exist_ok=True)
                    module.plot_named(names_map, path)
                c['path'] = str(path)
                swaps.append((stem, path))
    print(f"  step 2: {len(swaps)} customer chart(s) -> their {'REALNAMES' if real else 'twin'} variant: "
          f"{', '.join(stem for stem, _ in swaps) or 'none'}")

    # 3. Takeaways
    changed = fit_takeaways([s for sec in deck['sections'] for s in sec['slides']], font_file)
    print(f"  step 3: takeaways fitted to two lines ({Path(font_file).name}): {', '.join(changed) or 'no change'}")

    # Build, then check.
    out = OUT_REAL if real else Path(twin_dir) / TWIN_FILE
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.mkdtemp(prefix='followups_pptx_', dir=tmp_parent or names.DELIVERABLES))
    try:
        content = tmp / 'content.json'
        content.write_text(json.dumps(deck, ensure_ascii=False))
        built = tmp / out.name
        try:
            r = subprocess.run(['node', str(BUILDER), '--assets', str(assets), '--content', str(content),
                                '--out', str(built)], capture_output=True, text=True, check=False)
        except FileNotFoundError:
            fail('node not found: install Node and set NODE_PATH to the node_modules holding pptxgenjs')
        if r.returncode:
            fail('followups_pptx.js failed: ' + redact((r.stderr or r.stdout).strip()[-600:]))
        print('  ' + re.sub(r'wrote \S*?([^/\\]+\.pptx)', r'built \1', redact(r.stdout.strip())))
        n_slides = 2 + sum(len(sec['slides']) for sec in deck['sections']) + 1
        problems = check(built, n_slides, names_map, counter['n'], swaps, real)
        if problems:
            fail('checks failed: ' + '; '.join(problems))
        shutil.move(str(built), str(out))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print(f'  wrote {out.relative_to(ROOT) if out.is_relative_to(ROOT) else out} ({out.stat().st_size / 1e6:.1f} MB)')


def check(pptx_file, n_slides, names_map, n_substituted, swaps, real):
    """The written deck's checks, printed as counts. Returns the failed checks' descriptions."""
    from pptx import Presentation
    prs = Presentation(str(pptx_file))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            texts += [str(d) for d in shape._element.xpath('.//p:cNvPr/@descr')]
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
    blob = '\n'.join(texts)
    with zipfile.ZipFile(pptx_file) as z:
        media = {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('ppt/media/')}
    results = [
        ('slides', len(prs.slides) == n_slides, f'{len(prs.slides)} slides ({n_slides} expected)'),
        ('privacy: pseudonyms', names.pseudonyms_left(blob) == 0,
         f'{names.pseudonyms_left(blob)} pseudonyms in {len(texts)} text blocks (slide text, notes, alt text)'),
        ('privacy: names', names.names_found(blob, names_map) == n_substituted,
         f'{names.names_found(blob, names_map)} names found, {n_substituted} substituted'),
        ('customer charts', all(sha256(p) in media for _, p in swaps)
         and not any(sha256(FIG / f'{stem}.png') in media for stem, _ in swaps),
         f'{sum(sha256(p) in media for _, p in swaps)} of {len(swaps)} variants embedded byte for byte, '
         f'{sum(sha256(FIG / f"{stem}.png") in media for stem, _ in swaps)} pseudonymous ones'),
    ]
    validator = find_validator()
    if validator is None:
        print("  validate.py: skipped (the pptx skill's scripts/office/validate.py was not found; set PPTX_VALIDATE)")
    else:
        r = subprocess.run([sys.executable, str(validator), str(pptx_file)], capture_output=True, text=True,
                           cwd=validator.parent, check=False)
        summary = [line for line in (r.stdout + r.stderr).splitlines() if 'PASSED' in line or 'FAILED' in line]
        results.append(('validate.py', r.returncode == 0 and any('PASSED' in s for s in summary),
                        summary[-1].strip() if summary else 'no summary line'))
    for name, ok, detail in results:
        print(f"  [{'ok' if ok else 'FAIL'}] {name}: {detail}")
    return [name for name, ok, _ in results if not ok]


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument('--assets', required=True,
                    help="folder with the deck's mit_logo.png, ipa_logo.png and sloan_photo.jpg")
    ap.add_argument('--twin', metavar='OUT_DIR', help='build the twin-name copy into OUT_DIR instead')
    ap.add_argument('--tmp', metavar='DIR', help='parent of the temporary folder (default deliverables/)')
    a = ap.parse_args()
    try:
        build(Path(a.assets), twin_dir=a.twin, tmp_parent=a.tmp)
    except Failed as err:
        print(f'FAILED: {err}')
        sys.exit(1)
    except Exception as err:  # anything else: report it with any name it might quote redacted, never a traceback
        real = names.real_names() or {}
        print(f'FAILED: unexpected {type(err).__name__}: {redactor(real, names.twin(real))(str(err))}')
        sys.exit(1)


if __name__ == '__main__':
    main()
